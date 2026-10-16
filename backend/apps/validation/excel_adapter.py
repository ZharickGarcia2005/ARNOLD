"""Adaptador seguro de libros XLSX al contrato puro del motor financiero."""

from __future__ import annotations

from collections import Counter
from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
from pathlib import Path
from typing import BinaryIO
import uuid
import zipfile

from django.conf import settings
from defusedxml import ElementTree as SafeElementTree
from defusedxml.common import DefusedXmlException
from openpyxl import load_workbook

from .contracts import CellKind, DatasetInput, DatasetRole, FieldCell, RecordInput
from .rules import SOURCE_PROFILES, make_unique_headers, normalize_header, recognize_schema


ADAPTER_VERSION = "2026.10.1"
_REQUIRED_PARTS = {"[Content_Types].xml", "xl/workbook.xml"}
_FORBIDDEN_PREFIXES = (
    "xl/activex/",
    "xl/embeddings/",
    "xl/externallinks/",
)
_FORBIDDEN_SUFFIXES = ("vbaproject.bin",)


class ExcelAdapterError(ValueError):
    """Error permanente y seguro para mostrar al usuario sin exponer datos."""

    def __init__(self, code: str, message: str):
        super().__init__(message)
        self.code = code


def _rewind(source: str | Path | BinaryIO) -> None:
    if hasattr(source, "seek"):
        source.seek(0)


def inspect_xlsx_archive(source: str | Path | BinaryIO) -> None:
    """Rechaza contenedores XLSX peligrosos antes de invocar openpyxl."""

    max_entries = int(getattr(settings, "ARNOLD_XLSX_MAX_ENTRIES", 10_000))
    max_uncompressed = int(
        getattr(settings, "ARNOLD_XLSX_MAX_UNCOMPRESSED_BYTES", 250 * 1024 * 1024)
    )
    max_ratio = int(getattr(settings, "ARNOLD_XLSX_MAX_COMPRESSION_RATIO", 200))

    try:
        _rewind(source)
        with zipfile.ZipFile(source) as archive:
            infos = archive.infolist()
            names = {item.filename.replace("\\", "/") for item in infos}
            canonical_names = [name.casefold() for name in names]
            if len(names) != len(infos) or len(set(canonical_names)) != len(infos):
                raise ExcelAdapterError(
                    "xlsx_duplicate_part",
                    "El libro contiene componentes internos duplicados.",
                )
            if not _REQUIRED_PARTS.issubset(names):
                raise ExcelAdapterError("xlsx_structure", "El archivo no es un XLSX válido.")
            if len(infos) > max_entries:
                raise ExcelAdapterError("xlsx_too_complex", "El libro contiene demasiados componentes.")

            total = 0
            for item in infos:
                name = item.filename.replace("\\", "/")
                normalized_name = name.casefold()
                parts = name.split("/")
                if "\x00" in name or name.startswith("/") or ".." in parts:
                    raise ExcelAdapterError("xlsx_unsafe_path", "El libro contiene rutas internas no seguras.")
                if item.flag_bits & 0x1:
                    raise ExcelAdapterError("xlsx_encrypted", "No se admiten libros cifrados.")
                if normalized_name.startswith(_FORBIDDEN_PREFIXES) or normalized_name.endswith(
                    _FORBIDDEN_SUFFIXES
                ):
                    raise ExcelAdapterError(
                        "xlsx_active_content",
                        "No se admiten macros, enlaces externos ni objetos incrustados.",
                    )
                total += item.file_size
                if total > max_uncompressed:
                    raise ExcelAdapterError("xlsx_expanded_too_large", "El libro excede el tamaño seguro.")
                if item.file_size > 1_048_576:
                    ratio = item.file_size / max(item.compress_size, 1)
                    if ratio > max_ratio:
                        raise ExcelAdapterError("xlsx_compression_ratio", "El libro tiene compresión anormal.")
            for item in infos:
                if not item.filename.lower().endswith(".rels"):
                    continue
                if item.file_size > 1024 * 1024:
                    raise ExcelAdapterError("xlsx_unsafe_links", "El libro contiene relaciones no seguras.")
                try:
                    root = SafeElementTree.fromstring(archive.read(item))
                except (DefusedXmlException, SafeElementTree.ParseError, ValueError) as exc:
                    raise ExcelAdapterError("xlsx_unsafe_xml", "El libro contiene XML no seguro.") from exc
                if any(
                    str(element.attrib.get("TargetMode", "")).casefold() == "external"
                    for element in root.iter()
                ):
                    raise ExcelAdapterError("xlsx_external_link", "No se admiten enlaces externos.")
    except ExcelAdapterError:
        raise
    except (zipfile.BadZipFile, OSError, ValueError) as exc:
        raise ExcelAdapterError("xlsx_invalid", "El archivo no es un XLSX válido.") from exc
    finally:
        _rewind(source)


def _safe_scalar(value: object) -> object:
    if value is None or isinstance(value, (str, int, float, Decimal, bool, date, datetime)):
        return value
    return str(value)


def _identifier_raw(cell) -> str | None:
    if cell.value is None:
        return None
    numeric_value = cell.value
    if isinstance(numeric_value, float) and numeric_value.is_integer():
        numeric_value = int(numeric_value)
    if isinstance(numeric_value, int) and not isinstance(numeric_value, bool):
        number_format = str(cell.number_format or "")
        if number_format and set(number_format) == {"0"}:
            return f"{numeric_value:0{len(number_format)}d}"
        return str(numeric_value)
    return str(cell.value)


def _field_cell(cell, normalized_name: str) -> FieldCell:
    value = _safe_scalar(cell.value)
    raw = _identifier_raw(cell) if normalized_name in {"CLAVE", "CIU", "CEDULA_RUC"} else (
        None if cell.value is None else str(cell.value)
    )
    if cell.data_type == "f":
        kind = CellKind.FORMULA
        formula = str(cell.value or "")
        value = formula if formula.startswith("=") else f"={formula}"
        raw = str(value)
    elif cell.data_type == "e":
        kind = CellKind.ERROR
    elif cell.value is None or (isinstance(cell.value, str) and not cell.value.strip()):
        kind = CellKind.BLANK
    else:
        kind = CellKind.VALUE
    return FieldCell(value=value, raw=raw, kind=kind)


def _read_source(source: str | Path | BinaryIO):
    if isinstance(source, (str, Path)):
        return source
    _rewind(source)
    return BytesIO(source.read())


def parse_xlsx(
    source: str | Path | BinaryIO,
    *,
    batch_id: str,
    role: DatasetRole | str,
    period: date | None,
    schema_name: str,
    sheet_name: str | None = None,
    header_row: int = 1,
    profile_id: str | None = None,
    parent_dataset_id: str | None = None,
) -> DatasetInput:
    """Convierte un XLSX en un DatasetInput conservando procedencia y celdas."""

    if header_row < 1:
        raise ExcelAdapterError("header_row", "La fila de encabezados debe ser mayor que cero.")
    inspect_xlsx_archive(source)
    workbook_source = _read_source(source)
    try:
        workbook = load_workbook(
            workbook_source,
            read_only=True,
            data_only=False,
            keep_links=False,
        )
    except Exception as exc:
        raise ExcelAdapterError("xlsx_read", "No se pudo leer el libro XLSX.") from exc

    try:
        if sheet_name:
            if sheet_name not in workbook.sheetnames:
                raise ExcelAdapterError("sheet_missing", "La hoja solicitada no existe.")
            sheet = workbook[sheet_name]
        else:
            candidates = [item for item in workbook.worksheets if item.max_row >= header_row]
            if len(candidates) != 1:
                raise ExcelAdapterError(
                    "sheet_ambiguous",
                    "Seleccione explícitamente una hoja cuando el libro contiene varias.",
                )
            sheet = candidates[0]

        max_rows = int(getattr(settings, "ARNOLD_XLSX_MAX_ROWS", 100_000))
        max_columns = int(getattr(settings, "ARNOLD_XLSX_MAX_COLUMNS", 100))
        max_cells = int(getattr(settings, "ARNOLD_XLSX_MAX_CELLS", 2_000_000))
        if (
            sheet.max_row is not None and sheet.max_row > max_rows
        ) or (
            sheet.max_column is not None and sheet.max_column > max_columns
        ):
            raise ExcelAdapterError("sheet_limits", "La hoja excede los límites permitidos.")

        header_cells = next(
            sheet.iter_rows(min_row=header_row, max_row=header_row),
            (),
        )
        if not header_cells:
            raise ExcelAdapterError("headers_missing", "No se encontró la fila de encabezados.")
        if len(header_cells) > max_columns:
            raise ExcelAdapterError("sheet_limits", "La hoja excede los límites permitidos.")
        original_headers = tuple("" if cell.value is None else str(cell.value) for cell in header_cells)
        recognition = recognize_schema(original_headers, schema_name or None)
        resolved_schema = schema_name or recognition.schema_name or ""
        occurrences = make_unique_headers(original_headers, resolved_schema)

        profile = SOURCE_PROFILES.get(profile_id or "")
        if profile_id and profile is None:
            raise ExcelAdapterError("profile_unknown", "El perfil de columnas no está registrado.")
        if profile and profile.schema_name != resolved_schema:
            raise ExcelAdapterError("profile_schema", "El perfil no corresponde al esquema seleccionado.")
        if profile and profile.period and profile.period != period:
            raise ExcelAdapterError("profile_period", "El perfil no corresponde al periodo indicado.")
        selected = dict(profile.selected_occurrences) if profile else {}

        counts = Counter(item.normalized for item in occurrences)
        for name, occurrence in selected.items():
            if counts[name] < occurrence:
                raise ExcelAdapterError("profile_occurrence", "El perfil solicita una columna inexistente.")

        records: list[RecordInput] = []
        cells_read = len(header_cells)
        dataset_id = str(batch_id)
        for row in sheet.iter_rows(min_row=header_row + 1):
            if row and row[0].row > max_rows:
                raise ExcelAdapterError("sheet_limits", "La hoja excede los límites permitidos.")
            if all(cell.value is None for cell in row):
                continue
            cells_read += len(row)
            if cells_read > max_cells:
                raise ExcelAdapterError("sheet_limits", "La hoja excede los límites permitidos.")
            fields: dict[str, FieldCell] = {}
            by_base: dict[str, list[FieldCell]] = {}
            for occurrence, cell in zip(occurrences, row):
                converted = _field_cell(cell, occurrence.normalized)
                storage_name = (
                    f"{occurrence.normalized}__{occurrence.occurrence}"
                    if counts[occurrence.normalized] > 1
                    else occurrence.normalized
                )
                fields[storage_name] = converted
                by_base.setdefault(occurrence.normalized, []).append(converted)
            for base, cells in by_base.items():
                chosen = selected.get(base, 1)
                fields[base] = cells[chosen - 1]
            record_id = uuid.uuid5(
                uuid.NAMESPACE_URL,
                f"arnold:{dataset_id}:{sheet.title}:{row[0].row}",
            )
            records.append(
                RecordInput(
                    dataset_id=dataset_id,
                    record_id=str(record_id),
                    source_row=row[0].row,
                    fields=fields,
                )
            )

        if not records:
            raise ExcelAdapterError(
                "rows_missing",
                "La hoja no contiene registros debajo de los encabezados.",
            )

        return DatasetInput(
            dataset_id=dataset_id,
            batch_id=dataset_id,
            role=DatasetRole(role),
            period=period,
            schema_name=resolved_schema,
            original_headers=original_headers,
            records=tuple(records),
            profile_id=profile_id,
            parent_dataset_id=parent_dataset_id,
            selected_occurrences=selected,
        )
    finally:
        workbook.close()
