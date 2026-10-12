"""Exportación segura de reportes de incidencias a XLSX en memoria."""

from __future__ import annotations

from collections import Counter
from dataclasses import dataclass
from datetime import datetime
from decimal import Decimal
from io import BytesIO
from math import isfinite
from typing import Any, Mapping

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .contracts import RuleResult, Severity, ValidationReport, canonical_json
from .rules import RULE_CATALOG


@dataclass(frozen=True, slots=True)
class ExportMetadata:
    report_id: str
    generated_at: datetime
    generated_by: str
    approved_by: str | None = None
    classification: str = "CONFIDENCIAL"

    def __post_init__(self) -> None:
        if not self.report_id.strip():
            raise ValueError("report_id es obligatorio")
        if not self.generated_by.strip():
            raise ValueError("generated_by es obligatorio")


_DANGEROUS_PREFIXES = ("=", "+", "-", "@")
_MAX_EXCEL_TEXT_LENGTH = 32_767
_MAX_ISSUES_PER_WORKBOOK = 100_000
_HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
_SUBHEADER_FILL = PatternFill("solid", fgColor="D9EAF7")
_WHITE_BOLD = Font(color="FFFFFF", bold=True)
_TITLE_FONT = Font(size=14, bold=True, color="1F1F1F")


def safe_excel_value(value: Any) -> Any:
    """Neutraliza inyección de fórmulas sin convertir números a texto."""

    if value is None or isinstance(value, (bool, int, datetime)):
        return value
    if isinstance(value, float):
        return value if isfinite(value) else str(value)
    if isinstance(value, Decimal):
        return value if value.is_finite() else str(value)
    text = str(value)
    if text.lstrip().startswith(_DANGEROUS_PREFIXES):
        text = "'" + text
    if len(text) > _MAX_EXCEL_TEXT_LENGTH:
        raise ValueError(
            "Un texto supera el límite seguro de 32.767 caracteres de Excel; "
            "la exportación se detuvo para no truncar evidencia."
        )
    return text


def _changes(issue: RuleResult) -> tuple[str, str]:
    changes = issue.evidence.get("cambios")
    if not isinstance(changes, (list, tuple)):
        return "", ""
    previous: list[str] = []
    current: list[str] = []
    for item in changes:
        if not isinstance(item, Mapping):
            continue
        field = str(item.get("campo", ""))
        previous.append(f"{field}: {item.get('anterior', '')}")
        current.append(f"{field}: {item.get('actual', '')}")
    return " | ".join(previous), " | ".join(current)


def _detail_row(report: ValidationReport, issue: RuleResult) -> list[Any]:
    subject = issue.subject
    previous, current = _changes(issue)
    related_rows = ", ".join(
        f"{item.dataset_id}:{item.source_row}" for item in issue.related
    )
    definition = RULE_CATALOG[issue.rule_code]
    return [
        subject.period.isoformat() if subject and subject.period else "",
        report.current_dataset_id,
        report.reference_dataset_id or "",
        subject.source_row if subject else None,
        subject.record_id if subject else "",
        subject.clave if subject else "",
        subject.ciu if subject else "",
        subject.cedula_ruc if subject else "",
        subject.nombres if subject else "",
        issue.rule_code.value,
        definition.title,
        issue.severity.value,
        "Sí" if issue.blocking else "No",
        "Sí" if issue.requires_review else "No",
        issue.field_name or "",
        issue.message,
        previous,
        current,
        related_rows,
        canonical_json(issue.evidence),
        issue.fingerprint,
        issue.engine_version,
    ]


def _style_header(cells) -> None:
    for cell in cells:
        cell.fill = _HEADER_FILL
        cell.font = _WHITE_BOLD
        cell.alignment = Alignment(horizontal="center", vertical="center")


def _fit_columns(sheet, widths: dict[int, int]) -> None:
    for index, width in widths.items():
        sheet.column_dimensions[get_column_letter(index)].width = width


def build_issues_workbook(
    report: ValidationReport,
    metadata: ExportMetadata,
) -> Workbook:
    if len(report.issues) > _MAX_ISSUES_PER_WORKBOOK:
        raise ValueError(
            "El reporte supera el límite seguro de 100.000 incidencias por XLSX; "
            "debe dividirse antes de exportarlo."
        )
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Resumen"
    details = workbook.create_sheet("Incidencias")

    summary.sheet_view.showGridLines = False
    details.sheet_view.showGridLines = False

    summary["A1"] = "Reporte de incidencias financieras"
    summary["A1"].font = _TITLE_FONT
    summary.merge_cells("A1:D1")

    metadata_rows = [
        ("Clasificación", metadata.classification),
        ("ID del reporte", metadata.report_id),
        ("Dataset actual", report.current_dataset_id),
        ("Dataset de referencia", report.reference_dataset_id or "No proporcionado"),
        ("Versión del motor", report.engine_version),
        ("Fecha de generación", metadata.generated_at),
        ("Generado por", metadata.generated_by),
        ("Aprobado por", metadata.approved_by or "Pendiente"),
        ("Incidencias", len(report.issues)),
        ("Incidencias bloqueantes", report.blocking_count),
    ]
    for row_index, (label, value) in enumerate(metadata_rows, start=3):
        summary.cell(row_index, 1, label).font = Font(bold=True)
        summary.cell(row_index, 2, safe_excel_value(value))

    start_row = 15
    summary.cell(start_row, 1, "Código")
    summary.cell(start_row, 2, "Regla")
    summary.cell(start_row, 3, "Severidad")
    summary.cell(start_row, 4, "Cantidad")
    _style_header(summary[start_row])

    counts = Counter((item.rule_code, item.severity) for item in report.issues)
    row = start_row + 1
    for (code, severity), count in sorted(
        counts.items(), key=lambda item: (item[0][0].value, item[0][1].value)
    ):
        summary.cell(row, 1, code.value)
        summary.cell(row, 2, RULE_CATALOG[code].title)
        summary.cell(row, 3, severity.value)
        summary.cell(row, 4, count)
        row += 1

    row += 1
    summary.cell(row, 1, "Reglas ejecutadas").font = Font(bold=True)
    summary.cell(
        row,
        2,
        safe_excel_value(", ".join(code.value for code in report.executed_rules)),
    )
    row += 1
    summary.cell(row, 1, "Reglas omitidas").font = Font(bold=True)
    skipped_text = (
        " | ".join(
            f"{code.value}: {reason}"
            for code, reason in sorted(
                report.skipped_rules.items(), key=lambda item: item[0].value
            )
        )
        or "Ninguna"
    )
    summary.cell(
        row,
        2,
        safe_excel_value(skipped_text),
    )
    summary.cell(row, 2).alignment = Alignment(wrap_text=True, vertical="top")
    _fit_columns(summary, {1: 24, 2: 70, 3: 16, 4: 12})
    summary.freeze_panes = "A15"

    headers = [
        "Periodo",
        "Dataset actual",
        "Dataset referencia",
        "Fila original",
        "ID interno",
        "CLAVE",
        "CIU",
        "Cédula/RUC",
        "Nombre",
        "Código",
        "Regla",
        "Severidad",
        "Bloqueante",
        "Requiere revisión",
        "Campo",
        "Mensaje",
        "Valor anterior",
        "Valor actual",
        "Filas relacionadas",
        "Evidencia JSON",
        "Huella",
        "Versión motor",
    ]
    details.append(headers)
    _style_header(details[1])
    for issue in report.issues:
        details.append([safe_excel_value(value) for value in _detail_row(report, issue)])

    details.freeze_panes = "A2"
    details.auto_filter.ref = details.dimensions
    identifier_columns = (5, 6, 7, 8, 10, 21, 22)
    if details.max_row >= 2:
        for column in identifier_columns:
            for cell in details.iter_cols(
                min_col=column,
                max_col=column,
                min_row=2,
                max_row=details.max_row,
            ):
                for item in cell:
                    item.number_format = "@"
        for row_cells in details.iter_rows(min_row=2):
            for cell in row_cells:
                cell.alignment = Alignment(vertical="top", wrap_text=False)
    _fit_columns(
        details,
        {
            1: 12,
            2: 20,
            3: 20,
            4: 12,
            5: 20,
            6: 16,
            7: 16,
            8: 18,
            9: 28,
            10: 14,
            11: 28,
            12: 12,
            13: 12,
            14: 18,
            15: 18,
            16: 50,
            17: 30,
            18: 30,
            19: 24,
            20: 70,
            21: 66,
            22: 16,
        },
    )
    details.row_dimensions[1].height = 28
    return workbook


def render_issues_xlsx(
    report: ValidationReport,
    metadata: ExportMetadata,
) -> bytes:
    workbook = build_issues_workbook(report, metadata)
    output = BytesIO()
    workbook.save(output)
    payload = output.getvalue()
    verify_xlsx(payload)
    return payload


def verify_xlsx(payload: bytes) -> None:
    """Validación mínima reutilizable por el futuro servicio de almacenamiento."""

    workbook = load_workbook(BytesIO(payload), data_only=False, read_only=False)
    if workbook.sheetnames != ["Resumen", "Incidencias"]:
        raise ValueError("El XLSX no contiene las hojas esperadas")
    if getattr(workbook, "_external_links", ()):
        raise ValueError("La exportación no puede contener vínculos externos")
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if cell.data_type == "f":
                    raise ValueError("La exportación no puede contener fórmulas")
                if cell.hyperlink is not None:
                    raise ValueError("La exportación no puede contener hipervínculos")
