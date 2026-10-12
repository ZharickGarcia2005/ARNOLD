from datetime import datetime
from io import BytesIO
from unittest.mock import patch

from django.test import SimpleTestCase
from openpyxl import Workbook, load_workbook

from apps.validation.contracts import (
    ENGINE_VERSION,
    RuleCode,
    RuleResult,
    Severity,
    SubjectRef,
    ValidationReport,
)
from apps.validation.exporters import (
    ExportMetadata,
    render_issues_xlsx,
    safe_excel_value,
    verify_xlsx,
)

from .factories import dataset, record


class ExporterTests(SimpleTestCase):
    databases = []

    def setUp(self):
        self.metadata = ExportMetadata(
            report_id="reporte-ficticio-001",
            generated_at=datetime(2026, 10, 9, 10, 30),
            generated_by="analista-1",
            approved_by="jefe-1",
        )

    def _report(self) -> ValidationReport:
        source = dataset(
            record(
                "fila-2",
                2,
                clave="000123",
                nombres="=HYPERLINK(\"https://example.invalid\")",
            )
        )
        subject = SubjectRef.from_record(source.records[0], source)
        critical = RuleResult(
            rule_code=RuleCode.KEY_001,
            severity=Severity.CRITICAL,
            blocking=True,
            requires_review=True,
            subject=subject,
            field_name="CLAVE",
            message="+texto que no debe convertirse en fórmula",
            evidence={"tipo": "clave_repetida", "filas": [2, 8]},
        )
        informational = RuleResult(
            rule_code=RuleCode.MONTH_001,
            severity=Severity.INFO,
            blocking=False,
            requires_review=False,
            subject=subject,
            message="Alta del periodo",
            evidence={"tipo": "alta", "implica_pago": False},
        )
        return ValidationReport(
            engine_version=ENGINE_VERSION,
            current_dataset_id="actual",
            reference_dataset_id="referencia",
            issues=(critical, informational),
            executed_rules=(RuleCode.KEY_001, RuleCode.MONTH_001),
        )

    @staticmethod
    def _open(payload: bytes):
        return load_workbook(BytesIO(payload), data_only=False)

    def test_export_contains_only_summary_and_issue_sheets(self):
        payload = render_issues_xlsx(self._report(), self.metadata)
        workbook = self._open(payload)

        self.assertTrue(payload.startswith(b"PK"))
        self.assertEqual(workbook.sheetnames, ["Resumen", "Incidencias"])
        verify_xlsx(payload)

    def test_summary_counts_reconcile_with_detail_rows(self):
        workbook = self._open(render_issues_xlsx(self._report(), self.metadata))
        summary = workbook["Resumen"]
        details = workbook["Incidencias"]
        counts = {}
        for row in range(16, summary.max_row + 1):
            code = summary.cell(row, 1).value
            quantity = summary.cell(row, 4).value
            if isinstance(code, str) and code in {item.value for item in RuleCode}:
                counts[code] = quantity

        detail_codes = [details.cell(row, 10).value for row in range(2, details.max_row + 1)]

        self.assertEqual(counts, {"KEY-001": 1, "MONTH-001": 1})
        self.assertEqual(sum(counts.values()), len(detail_codes))
        self.assertEqual(sorted(detail_codes), ["KEY-001", "MONTH-001"])

    def test_identifiers_with_leading_zeroes_are_text(self):
        workbook = self._open(render_issues_xlsx(self._report(), self.metadata))
        details = workbook["Incidencias"]

        self.assertEqual(details["F2"].value, "000123")
        self.assertEqual(details["F2"].number_format, "@")
        self.assertEqual(details["G2"].value, "000987")
        self.assertEqual(details["G2"].number_format, "@")

    def test_export_neutralizes_formula_like_text_and_contains_no_formulas(self):
        workbook = self._open(render_issues_xlsx(self._report(), self.metadata))
        details = workbook["Incidencias"]

        self.assertTrue(details["I2"].value.startswith("'="))
        self.assertTrue(details["P2"].value.startswith("'+"))
        for sheet in workbook.worksheets:
            for row in sheet.iter_rows():
                for item in row:
                    self.assertNotEqual(item.data_type, "f")

    def test_formula_like_text_after_whitespace_is_also_neutralized(self):
        self.assertTrue(safe_excel_value("\t=1+1").startswith("'"))
        self.assertTrue(safe_excel_value("  +CMD").startswith("'"))

    def test_export_includes_traceability_metadata(self):
        workbook = self._open(render_issues_xlsx(self._report(), self.metadata))
        summary = workbook["Resumen"]
        values = {
            summary.cell(row, 1).value: summary.cell(row, 2).value
            for row in range(3, 13)
        }

        self.assertEqual(values["ID del reporte"], "reporte-ficticio-001")
        self.assertEqual(values["Dataset actual"], "actual")
        self.assertEqual(values["Dataset de referencia"], "referencia")
        self.assertEqual(values["Versión del motor"], ENGINE_VERSION)
        self.assertEqual(values["Generado por"], "analista-1")
        self.assertEqual(values["Aprobado por"], "jefe-1")

    def test_evidence_json_is_stable_and_fingerprint_is_exported(self):
        report = self._report()
        workbook = self._open(render_issues_xlsx(report, self.metadata))
        details = workbook["Incidencias"]

        self.assertEqual(
            details["T2"].value,
            '{"filas":[2,8],"tipo":"clave_repetida"}',
        )
        self.assertEqual(details["U2"].value, report.issues[0].fingerprint)

    def test_empty_report_still_produces_a_valid_xlsx(self):
        report = ValidationReport(
            engine_version=ENGINE_VERSION,
            current_dataset_id="actual-vacio",
            reference_dataset_id=None,
            issues=(),
            executed_rules=(),
        )

        payload = render_issues_xlsx(report, self.metadata)
        workbook = self._open(payload)

        populated_detail_rows = [
            row
            for row in workbook["Incidencias"].iter_rows(min_row=2, values_only=True)
            if any(value is not None for value in row)
        ]
        self.assertEqual(populated_detail_rows, [])
        self.assertEqual(workbook["Resumen"]["B11"].value, 0)
        self.assertEqual(workbook["Resumen"]["B12"].value, 0)
        verify_xlsx(payload)

    def test_export_rejects_more_issues_than_the_safe_workbook_limit(self):
        with patch("apps.validation.exporters._MAX_ISSUES_PER_WORKBOOK", 1):
            with self.assertRaisesRegex(ValueError, "límite seguro"):
                render_issues_xlsx(self._report(), self.metadata)

    def test_public_renderer_verifies_payload_before_returning_it(self):
        with patch("apps.validation.exporters.verify_xlsx") as verifier:
            payload = render_issues_xlsx(self._report(), self.metadata)

        verifier.assert_called_once_with(payload)

    def test_export_fails_closed_instead_of_truncating_long_evidence(self):
        source = dataset(record("fila-2", 2))
        subject = SubjectRef.from_record(source.records[0], source)
        issue = RuleResult(
            rule_code=RuleCode.KEY_001,
            severity=Severity.CRITICAL,
            blocking=True,
            requires_review=True,
            subject=subject,
            evidence={"detalle": "x" * 33_000},
        )
        report = ValidationReport(
            engine_version=ENGINE_VERSION,
            current_dataset_id="actual",
            reference_dataset_id=None,
            issues=(issue,),
            executed_rules=(RuleCode.KEY_001,),
        )

        with self.assertRaisesRegex(ValueError, "no truncar evidencia"):
            render_issues_xlsx(report, self.metadata)

    def test_export_fails_closed_for_long_summary_text(self):
        report = ValidationReport(
            engine_version=ENGINE_VERSION,
            current_dataset_id="actual",
            reference_dataset_id=None,
            issues=(),
            executed_rules=(),
            skipped_rules={RuleCode.KEY_002: "x" * 33_000},
        )

        with self.assertRaisesRegex(ValueError, "no truncar evidencia"):
            render_issues_xlsx(report, self.metadata)

    def test_verifier_rejects_hyperlinks(self):
        workbook = Workbook()
        workbook.active.title = "Resumen"
        workbook.create_sheet("Incidencias")
        workbook["Resumen"]["A1"] = "Enlace"
        workbook["Resumen"]["A1"].hyperlink = "https://example.invalid"
        output = BytesIO()
        workbook.save(output)

        with self.assertRaisesRegex(ValueError, "hipervínculos"):
            verify_xlsx(output.getvalue())
