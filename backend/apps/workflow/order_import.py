"""Importación explícita del libro operativo, sin sustituir carteras mensuales.

Una sola cartera conserva las doce hojas confirmadas. source_row es un ordinal
interno único; hoja, fila Excel y encabezados literales se autentican en el RAW.
No se ejecutan fórmulas ni se utilizan sus resultados cacheados como aprobados.
"""
from dataclasses import replace
from hashlib import sha256
from pathlib import Path
import os
import tempfile
import uuid

from django.conf import settings
from django.db import transaction
from django.utils import timezone
from openpyxl import load_workbook

from apps.audit.services import append_audit_event
from apps.batches.crypto import encrypt_json
from apps.batches.models import Batch
from apps.batches.services import UploadError, _scan, batch_file_path
from apps.validation.contracts import DatasetInput, DatasetRole, ENGINE_VERSION, RuleCode, ValidationInput
from apps.validation.digests import validation_input_digest
from apps.validation.excel_adapter import ADAPTER_VERSION, inspect_xlsx_archive, parse_xlsx
from apps.validation.models import ValidationRun
from apps.validation.rules import make_unique_headers, validate
from apps.validation.services import _persist_dataset, _persist_report, _record_payload, _record_aad
from .columns import SHEETS

IMPORT_PROFILE = "order_workbook_operational_v1"
SOURCE_TYPE = "orden_preliminar"


def read_order_book(path, batch_id):
    """Preflight sin escrituras ni datos personales en el resumen."""
    path = Path(path)
    inspect_xlsx_archive(path)
    book = load_workbook(path, read_only=True, data_only=False, keep_links=False)
    try:
        required = [definition["name"] for definition in SHEETS.values()]
        if set(required) - set(book.sheetnames):
            raise UploadError("missing_order_sheets", "Faltan hojas operativas confirmadas en el libro.")
        ignored = [name for name in book.sheetnames if name not in required]
        empty = set()
        for name in required:
            sheet = book[name]
            if ((sheet.max_row or 0) > settings.ARNOLD_XLSX_MAX_ROWS
                    or (sheet.max_column or 0) > settings.ARNOLD_XLSX_MAX_COLUMNS
                    or (sheet.max_row or 0) * (sheet.max_column or 0) > settings.ARNOLD_XLSX_MAX_CELLS):
                raise UploadError("sheet_limits", "Una hoja operativa supera los límites seguros.")
            if not any(cell.value is not None for row in sheet.iter_rows(min_row=2) for cell in row):
                empty.add(name)
    finally:
        book.close()
    rows, origins, header_union, seen = [], {}, [], set()
    counts, formulas, errors = {}, 0, 0
    for code, definition in SHEETS.items():
        name = definition["name"]
        if name in empty:
            counts[name] = 0
            continue
        part = parse_xlsx(path, batch_id=str(batch_id), role=DatasetRole.CURRENT,
                          period=None, schema_name="ordenPreliminar", sheet_name=name)
        for occurrence in make_unique_headers(part.original_headers, "ordenPreliminar"):
            identity = (occurrence.normalized, occurrence.occurrence)
            if identity not in seen:
                seen.add(identity)
                header_union.append(occurrence.original)
        counts[name] = len(part.records)
        for original in part.records:
            ordinal = len(rows) + 1
            rows.append(replace(original, source_row=ordinal))
            origins[original.record_id] = {"sheet": name, "excel_row": original.source_row,
                                           "order": "" if code == "issues" else code,
                                           "headers": list(part.original_headers)}
            formulas += sum(cell.kind.value == "formula" for key, cell in original.fields.items() if "__" not in key)
            errors += sum(cell.kind.value == "error" for key, cell in original.fields.items() if "__" not in key)
    if not rows:
        raise UploadError("rows_missing", "Las hojas operativas no contienen registros.")
    if len(rows) > settings.ARNOLD_XLSX_MAX_ROWS:
        raise UploadError("row_limit", "El libro operativo supera el límite de registros.")
    dataset = DatasetInput(dataset_id=str(batch_id), batch_id=str(batch_id), role=DatasetRole.CURRENT,
                           period=None, schema_name="ordenPreliminar", original_headers=tuple(header_union),
                           records=tuple(rows))
    return dataset, origins, {"sheets": counts, "ignored_sheets": ignored,
                              "rows": len(rows), "formula_cells": formulas, "error_cells": errors}


def import_order_book(path, *, actor, dry_run=False):
    """Entrada administrativa con los mismos permisos, cuarentena, cifrado y motor.

    La copia privada se examina antes de la transacción. Un fallo revierte todo;
    repetir el mismo archivo/área reutiliza el lote en vez de duplicar personas.
    """
    if not actor.is_active or not actor.has_perm("batches.import_batch"):
        raise UploadError("permission_denied", "El usuario no tiene permiso para importar.", status=403)
    area = actor.area.strip()
    if not actor.is_superuser and not area:
        raise UploadError("area_required", "El usuario no tiene un área asignada.", status=403)
    path = Path(path)
    if path.suffix.lower() != ".xlsx":
        raise UploadError("file_type", "Se requiere un archivo XLSX.")
    incoming = Path(settings.ARNOLD_FILE_ROOT) / "quarantine" / ".incoming"
    incoming.mkdir(parents=True, exist_ok=True, mode=0o700)
    temporary, destination = None, None
    try:
        digest, size = sha256(), 0
        with tempfile.NamedTemporaryFile(dir=incoming, suffix=".part", delete=False) as target:
            temporary = Path(target.name)
            with path.open("rb") as source:
                for chunk in iter(lambda: source.read(1048576), b""):
                    size += len(chunk)
                    if size > settings.ARNOLD_MAX_UPLOAD_BYTES:
                        raise UploadError("file_too_large", "El libro supera el límite permitido.")
                    target.write(chunk)
                    digest.update(chunk)
        os.chmod(temporary, 0o600)
        fingerprint = digest.hexdigest()
        # openpyxl requiere una extensión conocida, no el sufijo .part.
        source_copy = temporary.with_suffix(".xlsx")
        temporary.rename(source_copy)
        temporary = source_copy
        inspect_xlsx_archive(temporary)
        scan_status = _scan(temporary)
        existing = Batch.objects.filter(sha256=fingerprint, owner_area=area,
                                         source_type=SOURCE_TYPE, profile_id=IMPORT_PROFILE).first()
        if existing:
            return existing, {"reused": True, "rows": existing.row_count}
        batch_id = uuid.uuid4()
        dataset, origins, stats = read_order_book(temporary, batch_id)
        stats.update(sha256=fingerprint, reused=False)
        if dry_run:
            return None, stats
        report = validate(ValidationInput(current=dataset))
        with transaction.atomic():
            # Serializa las importaciones de la misma área, incluida la ausencia
            # de periodo: SQLite permite varios NULL en el índice de recepción.
            type(actor).objects.filter(area=actor.area).order_by("pk").select_for_update().first()
            existing = Batch.objects.filter(sha256=fingerprint, owner_area=area,
                                             source_type=SOURCE_TYPE, profile_id=IMPORT_PROFILE).first()
            if existing:
                return existing, {"reused": True, "rows": existing.row_count}
            batch = Batch.objects.create(id=batch_id, display_name=path.name, sha256=fingerprint,
                                         source_type=SOURCE_TYPE, period=None, status=Batch.Status.PROCESSING,
                                         file_size=size, content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                                         owner_area=area, sheet_name="Libro operativo 030–040 + INCONSISTENCIAS", header_row=1,
                                         schema_name="ordenPreliminar", profile_id=IMPORT_PROFILE,
                                         adapter_version=ADAPTER_VERSION, scan_status=scan_status, uploaded_by=actor)
            destination = batch_file_path(batch)
            os.replace(temporary, destination)
            temporary = None
            _persist_dataset(batch, dataset)
            originals = {item.record_id: item for item in dataset.records}
            pending = []
            for record in batch.records.iterator(chunk_size=500):
                raw = {**_record_payload(originals[str(record.pk)]), "origin": origins[str(record.pk)]}
                record.data_ciphertext = encrypt_json(raw, associated_data=_record_aad(batch.pk, record.pk, record.source_row))
                pending.append(record)
                if len(pending) == 500:
                    batch.records.model.objects.bulk_update(pending, ("data_ciphertext",)); pending.clear()
            if pending:
                batch.records.model.objects.bulk_update(pending, ("data_ciphertext",))
            run = ValidationRun.objects.create(current_batch=batch, created_by=actor,
                    input_digest=validation_input_digest(current_batch=batch), engine_version=ENGINE_VERSION,
                    adapter_version=ADAPTER_VERSION, enabled_rules=[rule.value for rule in RuleCode],
                    status=ValidationRun.Status.RUNNING, task_id=f"order-import:{uuid.uuid4()}",
                    attempt_count=1, started_at=timezone.now())
            append_audit_event(actor=actor, action="batch.order_workbook_received", object_type="Batch", object_id=batch.pk,
                               metadata={"sha256": fingerprint, "sheets": stats["sheets"], "ignored_sheets": stats["ignored_sheets"],
                                         "rows": stats["rows"], "scan_status": scan_status, "profile": IMPORT_PROFILE})
            _persist_report(run, report)
        destination = None  # confirmado; jamás borrar un original persistido
        batch.refresh_from_db()
        return batch, stats
    finally:
        if temporary is not None:
            temporary.unlink(missing_ok=True)
        if destination is not None:
            destination.unlink(missing_ok=True)  # sólo la copia nueva de una transacción fallida
