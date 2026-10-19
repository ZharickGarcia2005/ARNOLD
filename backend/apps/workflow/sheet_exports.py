"""Las dos descargas usan exactamente el contrato visible de la Worksheet."""
from copy import copy
from datetime import date, datetime, time
from decimal import Decimal
from hashlib import sha256
import json
from tempfile import SpooledTemporaryFile

from django.db import transaction
from django.db.models import Q
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.cell.cell import ERROR_CODES
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from apps.audit.services import append_audit_event
from apps.batches.crypto import decrypt_json
from apps.validation.exporters import safe_excel_value
from .columns import CONTRACT_VERSION, SHEETS, cell_value, columns_for, source_values, original_payload
from .order_projection import routes, selected_ids, decorate
from .models import ChangeRequest, Worksheet, WorksheetEdit
from .worksheet import MONEY, SheetError, filter_rows, row_payload
from .presentation import effective_style, read_document

export_columns = columns_for


def typed_value(field, value):
    if value not in (None, ""):
        if field in MONEY or field in ("EMISION", "INTERES", "COACTIVA", "RECARGO", "SUBTOTAL", "TOTAL", "SOURCE_TOTAL"):
            try:
                number = Decimal(str(value))
                if number.is_finite() and abs(number) <= Decimal("999999999999.99"):
                    value = number
            except ArithmeticError:
                pass
        elif field in ("MESES_DEUDA", "NO", "#"):
            try:
                number = Decimal(str(value))
                if number.is_finite() and 0 <= number <= 999999 and number == number.to_integral_value(): value = int(number)
            except ArithmeticError:
                pass
        elif field in ("FECHA", "FECHA_VENCIMIENTO") and isinstance(value, str):
            try: value = datetime.fromisoformat(value).date() if "T" in value or " " in value else date.fromisoformat(value)
            except ValueError: pass
        elif field == "HORA" and isinstance(value, str):
            try: value = time.fromisoformat(value)
            except ValueError: pass
    if isinstance(value, str) and value in ERROR_CODES:
        return "'" + value  # evidencia del error, no una celda de error ejecutable
    return value if isinstance(value, (date, time)) else safe_excel_value(value)


def _visual(cell, style):
    if style.get("background"): cell.fill = PatternFill("solid", fgColor=style["background"][1:])
    font = copy(cell.font)
    for key in ("bold", "italic"):
        if key in style: setattr(font, key, style[key])
    if "underline" in style: font.underline = "single" if style["underline"] else None
    if style.get("color"): font.color = style["color"][1:]
    if style.get("size"): font.sz = style["size"] * .75
    cell.font = font
    alignment = copy(cell.alignment)
    if style.get("align"): alignment.horizontal = style["align"]
    if "wrap" in style: alignment.wrap_text = style["wrap"]
    cell.alignment = alignment


class Trace:
    def __init__(self, book): self.book, self.sheet, self.count, self.part = book, None, 0, 0
    def append(self, key, value):
        if self.sheet is None or self.count >= 1048576:
            self.part += 1; self.count = 0
            self.sheet = self.book.create_sheet("Trazabilidad" if self.part == 1 else f"Trazabilidad {self.part}")
            self.sheet.column_dimensions["A"].width = 55; self.sheet.column_dimensions["B"].width = 110
        if isinstance(value, (dict, list)): value = json.dumps(value, ensure_ascii=False, default=str)
        self.sheet.append([safe_excel_value(key), typed_value("", value)]); self.count += 1


@transaction.atomic
def export_sheet(batch, actor, scope, search="", expected_revision=None):
    if scope not in ("clean", "problem"):
        raise SheetError("invalid_scope", "Seleccione el lote principal o las filas con problemas.")
    sheet = Worksheet.objects.select_for_update().get(batch=batch)
    presentation = read_document(sheet, actor)
    if expected_revision != sheet.revision:
        raise SheetError("version_conflict", "La hoja cambió. Actualice antes de descargar.", 409)
    # Las descargas siempre son completas, incluso si el cliente manda búsqueda.
    selected = selected_ids(sheet, scope)
    selected_set = {str(pk) for pk in selected}
    mapping = routes(sheet)
    query = filter_rows(sheet, "all")
    count = len(selected)
    if not count: raise SheetError("empty_export", "No hay registros en este conjunto. No se generó un Excel vacío.", 422)
    if count > 100000: raise SheetError("export_limit", "Divida el lote antes de exportar más de 100.000 filas.")
    book = Workbook(write_only=True)
    stream = SpooledTemporaryFile(max_size=8 * 1024 * 1024, mode="w+b")
    try:
        destinations, row_numbers = {}, {}
        border = Side(style="thin", color="B8C9DB")
        for code, definition in SHEETS.items():
            if (scope == "problem") != (code == "issues"): continue
            data = book.create_sheet(definition["name"])
            data.freeze_panes = "A2"; data.row_dimensions[1].height = 36
            data.sheet_view.showGridLines = False
            columns = columns_for(batch, code)
            header = []
            for index, column in enumerate(columns, 1):
                width = presentation["columns"].get(column.field, {}).get("width", column.width)
                dimension = data.column_dimensions[get_column_letter(index)]
                dimension.width = round((width - 5) / 7, 2)
                dimension.hidden = column.hidden; dimension.outlineLevel = column.outline
                cell = WriteOnlyCell(data, safe_excel_value(column.label))
                cell.font = Font(bold=True, color="243E54", name="Arial", size=10)
                cell.fill = PatternFill("solid", fgColor=definition["header_fill"][1:])
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                cell.border = Border(bottom=border, right=border)
                header.append(cell)
            data.append(header)
            destinations[code] = (data, columns); row_numbers[code] = 1
        trace = Trace(book)
        run = batch.validation_runs.order_by("-created_at").first()
        for key, value in (("Clasificación", "CONFIDENCIAL"), ("Lote", str(batch.pk)), ("Archivo original", batch.display_name),
                           ("Periodo", str(batch.period)), ("SHA256 original", batch.sha256), ("Hoja original", batch.sheet_name),
                           ("Revisión", sheet.revision), ("Conjunto", scope), ("Filas", count), ("Usuario", actor.username),
                           ("Generado", timezone.now().isoformat()), ("Contrato columnas", CONTRACT_VERSION),
                           ("Aprobación", "Exportación preliminar. No certifica pago, cancelación ni aprobación global del lote."),
                           ("Búsqueda aplicada", "No"), ("Formato personal", f"Usuario {actor.username}. Versión {presentation['version']}"),
                           ("Adaptador", run.adapter_version if run else batch.adapter_version), ("Motor", run.engine_version if run else "No informado"),
                           ("Reglas ejecutadas", run.executed_rules if run else []), ("Reglas omitidas", run.skipped_rules if run else {})):
            trace.append(key, value)
        for code, (data, columns) in destinations.items():
            for index, column in enumerate(columns): trace.append(f"Columna {code}/{index + 1}", column.payload(index))
        written, legacy = 0, 0
        for row in query.iterator(chunk_size=500):
            if str(row.pk) not in selected_set: continue
            code = "issues" if scope == "problem" else mapping[str(row.pk)][0]
            data, columns = destinations[code]
            row_numbers[code] += 1; index = row_numbers[code]
            payload = decorate(row_payload(row), mapping[str(row.pk)])
            payload["source_values"] = source_values(row, columns)
            height = presentation["heights"].get(str(row.pk))
            if height: data.row_dimensions[index].height = height * .75
            cells = []
            for column in columns:
                value = cell_value(payload, column)
                value = typed_value(column.field if column.working or column.control else column.source_field, value)
                cell = WriteOnlyCell(data, value)
                cell.number_format = column.number_format
                cell.font = Font(name="Arial", size=10, color="243E54")
                cell.fill = PatternFill("solid", fgColor="FFFFFF" if index % 2 == 0 else SHEETS[code]["stripe_fill"][1:])
                cell.border = Border(bottom=border, right=border)
                _visual(cell, effective_style(presentation, row.pk, column.field))
                cells.append(cell)
            data.append(cells); written += 1
            legacy += payload["identifier_validation"] != "executed"
            trace.append(f"Record {row.pk}", {"source_row": row.record.source_row, "worksheet_version": row.version,
                                               "source_origin": original_payload(row).get("origin", {}),
                                               "excel_row": index, "destination_sheet": data.title, "order": payload["order"], "identifier_validation": payload["identifier_validation"],
                                               "updated_by": row.updated_by_id, "quality": row.quality, "issues": payload["errors"],
                                               "missing_presentation_values": [col.label for col in columns if col.working and cell_value(payload, col) in ("", None)]})
        if written != count: raise SheetError("incomplete_export", "La cantidad de filas cambió. Actualice y reintente.", 409)
        for code, (data, columns) in destinations.items():
            data.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{row_numbers[code]}"
            trace.append(f"Filas {code}", row_numbers[code] - 1)
        trace.append("Filas con esquema histórico conservado", legacy)
        ids = selected
        if run:
            issues = run.issues.filter(Q(record_id__in=ids) | Q(related_records__in=ids)
                                      | Q(record__isnull=True, subject_batch=batch)).distinct().prefetch_related("related_records")
            for issue in issues.iterator(chunk_size=500):
                trace.append(f"Incidencia {issue.pk}", {"record": str(issue.record_id) if issue.record_id else None,
                             "related_records": [str(record.pk) for record in issue.related_records.all()],
                             "rule": issue.rule_code, "field": issue.field_name, "severity": issue.severity,
                             "blocking": issue.blocking, "requires_review": issue.requires_review,
                             "review_status": issue.review_status, "fingerprint": issue.fingerprint,
                             "evidence_sha256": issue.evidence_sha256, "engine_version": issue.engine_version})
        for edit in WorksheetEdit.objects.filter(row_id__in=ids).select_related("actor").iterator():
            detail = decrypt_json(edit.detail_ciphertext, associated_data=f"worksheet-edit:{edit.pk}")
            for field in detail["fields"]:
                trace.append(f"Cambio {edit.pk} / {field}", {"record": str(edit.row_id), "version": edit.version, "user": edit.actor.username,
                            "date": edit.created_at.isoformat(), "reason": detail["reason"], "old": detail["old"].get(field), "new": detail["new"].get(field)})
        for change in ChangeRequest.objects.filter(record_id__in=ids, field_name="__worksheet__").select_related("requested_by", "decided_by").iterator():
            old = decrypt_json(change.old_value_ciphertext, associated_data=f"sheet-change-old:{change.record_id}")
            new = decrypt_json(change.new_value_ciphertext, associated_data=f"sheet-change-new:{change.record_id}")
            for field, value in new.items():
                trace.append(f"Solicitud {change.pk} / {field}", {"record": str(change.record_id), "status": change.status, "requester": change.requested_by.username,
                             "decider": change.decided_by.username if change.decided_by else None, "expected_version": change.expected_version,
                             "requested": change.requested_at.isoformat(), "decided": change.decided_at.isoformat() if change.decided_at else None,
                             "reason": change.reason, "old": old.get(field), "proposed": value})
        book.save(stream); stream.seek(0)
        digest = sha256()
        for chunk in iter(lambda: stream.read(1048576), b""): digest.update(chunk)
        stream.seek(0)
        stream.arnold_export_rows, stream.arnold_export_revision = count, sheet.revision
        append_audit_event(actor=actor, action="worksheet.exported", object_type="Batch", object_id=batch.pk,
                           metadata={"scope": scope, "rows": count, "revision": sheet.revision, "sha256": digest.hexdigest(), "column_contract": CONTRACT_VERSION})
        return stream, f"ARNOLD_{scope}_{batch.period}_r{sheet.revision}.xlsx"
    except Exception:
        stream.close(); raise
    finally:
        book.close()
