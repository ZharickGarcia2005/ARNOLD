"""Contrato exacto por orden, partición segura y descarga multihoja completa."""
from datetime import datetime
from decimal import Decimal
from django.test import TestCase
from openpyxl.utils import get_column_letter
from . import test_column_contract as fixtures
from .columns import SHEETS
from .worksheet import decide_change, save_changes

class OrderPresentationTests(TestCase):
    setUp = fixtures.ColumnContractTests.setUp
    add_record = fixtures.ColumnContractTests.add_record
    prepare = fixtures.ColumnContractTests.prepare
    post = fixtures.ColumnContractTests.post
    download = fixtures.ColumnContractTests.download

    def test_exact_tabs_headers_hidden_widths_and_distinct_titles(self):
        for code in list(SHEETS)[:-1]: self.add_record(f"KEY-{code}", ORDEN_PRELIMINAR=code, TITULO="0007", TITULOS="0008")
        self.prepare()
        expected_names = ["030-SAN CARLOS (LISTA)", "031-(24 DE MAYO) LISTA", "032-SAN CRISTOBAL(LISTA)", "033-GUAYACAN(LISTA)",
                          "034-GUAYACAN(LISTA)", "035-NICOLAS ALTA(LISTA)", "036-RUTA DEL RIO(LISTA)", "037-VIVA ALFARO(LISTA)",
                          "038- VENUS", "039-VENUS", "040-VENUS", "INCONSISTENCIAS"]
        for code, template in SHEETS.items():
            data = self.post("query", {"order": code}).json()
            self.assertEqual([tab["name"] for tab in data["tabs"]], expected_names)
            actual = data["columns"][:len(template["columns"])]
            self.assertEqual([col["label"] for col in actual], [col["label"] for col in template["columns"]])
            self.assertEqual([col["hidden"] for col in actual], [col["hidden"] for col in template["columns"]])
            self.assertEqual([col["width"] for col in actual], [max(8, round(col["width"] * 7 + 5)) for col in template["columns"]])
            self.assertNotIn("FILA_ORIGEN", [col["label"] for col in actual])
            if code != "issues":
                self.assertEqual(len(data["rows"]), 1); self.assertEqual(data["rows"][0]["order"], code)
                self.assertEqual((data["rows"][0]["values"]["TITULO"], data["rows"][0]["values"]["TITULOS"]), ("0007", "0008"))
        visible = [col["label"].strip() for col in SHEETS["030"]["columns"] if not col["hidden"] and col["label"]]
        self.assertEqual(visible[:3], ["EXP", "PREDIO MUNICIPAL", "CONTRIBUYENTE"])
        self.assertEqual(SHEETS["030"]["columns"][21]["label"], "EMISION")
        self.assertEqual(SHEETS["036"]["columns"][22]["label"], "EMISION")
        self.assertEqual([col["label"].strip() for col in SHEETS["038"]["columns"]][-3:], ["UNIDAD", "FE", "REGISTRADO"])


    def test_ambiguous_missing_orders_go_to_issues_without_reclassifying_history(self):
        self.add_record("A", ORDEN_PRELIMINAR="", PARROQUIA="VENUS", OP="030")
        self.add_record("B", ORDEN_PRELIMINAR="999", PARROQUIA="GUAYACAN")
        sheet = self.prepare()
        snapshots = [(str(row.pk), bytes(row.data_ciphertext), row.quality, row.version) for row in sheet.rows.all()]
        data = self.post("query", {"order": "issues"}).json()
        self.assertEqual(data["summary"]["problem"], 2)
        self.assertEqual(data["validation_summary"]["clean"], 2)
        self.assertTrue(all(any(issue["code"] == "ORDER_MAPPING" for issue in row["errors"]) for row in data["rows"]))
        self.assertEqual(snapshots, [(str(row.pk), bytes(row.data_ciphertext), row.quality, row.version) for row in sheet.rows.all()])


    def test_approved_order_correction_moves_view_and_preserves_original(self):
        record = self.add_record("A", ORDEN_PRELIMINAR="")
        original = bytes(record.data_ciphertext); sheet = self.prepare()
        request = save_changes(self.batch, self.worker, [{"id": str(record.pk), "version": 1, "values": {"ORDEN_PRELIMINAR": "038"}}], "Orden confirmada sintética")
        self.assertEqual(self.post("query", {"order": "issues"}).json()["filtered_total"], 1)
        decide_change(self.batch, self.boss, request["pending"][0], True)
        self.assertEqual(self.post("query", {"order": "038"}).json()["filtered_total"], 1)
        record.refresh_from_db(); self.assertEqual(bytes(record.data_ciphertext), original)


    def test_two_downloads_keep_all_order_tabs_values_formats_and_counts(self):
        for number in range(105): self.add_record(f"030-{number}", EXP=f"EXP-{number}", FECHA="2026-10-01", ORDEN_PRELIMINAR="030")
        for code in list(SHEETS)[1:-1]: self.add_record(f"KEY-{code}", ORDEN_PRELIMINAR=code, UNIDAD="U-SINTETICA", FE="Dato FE", REGISTRADO="SI", SOURCE_TOTAL="19.00")
        self.add_record("BAD", ORDEN_PRELIMINAR="", NOMBRES="=1+1")
        self.prepare()
        ids = []
        for scope, count in (("clean", 115), ("problem", 1)):
            book, _ = self.download(scope, search="030-0", page=1, order="030")
            names = [template["name"] for code, template in SHEETS.items() if (code == "issues") == (scope == "problem")]
            self.assertEqual(book.sheetnames, names + ["Trazabilidad"])
            totals, rows_count = Decimal(0), 0
            for name in names:
                template_id = next(code for code, template in SHEETS.items() if template["name"] == name)
                expected = self.post("query", {"order": template_id}).json()["columns"]
                data = book[name]
                self.assertEqual([cell.value or "" for cell in data[1]], [col["label"] for col in expected])
                for index, column in enumerate(expected, 1):
                    self.assertEqual(data.column_dimensions[get_column_letter(index)].hidden, column["hidden"])
                    if data.max_row > 1: self.assertEqual(data.cell(2, index).number_format, column["number_format"])
                rows_count += data.max_row - 1
                for row in data.iter_rows(min_row=2):
                    labels = {cell.value: index for index, cell in enumerate(data[1])}
                    self.assertIsInstance(row[labels["CLAVE"]].value, str)
                    totals += Decimal(str(row[labels["EMISION"]].value))
                    self.assertFalse(any(cell.data_type in ("f", "e") for cell in row))
                self.assertEqual(data.freeze_panes, "A2")
            self.assertEqual(rows_count, count)
            self.assertEqual(totals, Decimal(10) * count)
            meta = dict(book["Trazabilidad"].values)
            self.assertEqual(meta["Filas"], count)
            ids.append({key for key in meta if str(key).startswith("Record ")})
            if scope == "clean":
                self.assertIsInstance(book[names[0]]["AH2"].value, datetime)
                self.assertEqual(book["038- VENUS"]["AM2"].value, "U-SINTETICA")
                self.assertNotEqual(book[names[0]]["E2"].fill.fgColor.rgb, book[names[0]]["E3"].fill.fgColor.rgb)
        self.assertFalse(ids[0] & ids[1]); self.assertEqual(len(ids[0] | ids[1]), 116)

    def test_editable_additional_money_date_and_time_are_validated(self):
        record = self.add_record("A", SOURCE_TOTAL="19.00", FECHA="2026-10-01", HORA="09:30:00")
        sheet = self.prepare()
        save_changes(self.batch, self.boss, [{"id": str(record.pk), "version": 1,
                     "values": {"SOURCE_TOTAL": "no-numérico", "FECHA": "2026-02-31", "HORA": "29:00"}}], "Prueba de campos de orden")
        row = sheet.rows.get(); self.assertEqual(row.quality, "problem")
        self.assertTrue({"SOURCE_TOTAL", "FECHA", "HORA"}.issubset({error["field"] for error in row.errors}))
