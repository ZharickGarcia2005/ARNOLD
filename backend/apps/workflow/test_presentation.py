"""Presentación aislada: permisos, concurrencia, paginación e integridad XLSX."""

import base64
from io import BytesIO
import tempfile

from django.contrib.auth import get_user_model
from django.conf import settings
from django.contrib.auth.models import Permission
from django.test import Client, TestCase, override_settings
from openpyxl import load_workbook

from .presentation import location, read_document, save_presentation
from .sheet_exports import export_sheet
from . import test_worksheet as fixtures
from .worksheet import SheetError, values_of


@override_settings(ARNOLD_DATA_ENCRYPTION_KEY=base64.b64encode(b"t" * 32).decode(), ARNOLD_HMAC_KEY="k" * 32)
class PresentationTests(TestCase):
    add_record = fixtures.WorksheetTests.add_record
    prepare = fixtures.WorksheetTests.prepare
    post = fixtures.WorksheetTests.post

    def setUp(self):
        fixtures.WorksheetTests.setUp(self)
        folder = tempfile.TemporaryDirectory(prefix="test-presentation-", dir=settings.BASE_DIR)
        self.addCleanup(folder.cleanup)
        self.root = override_settings(ARNOLD_FILE_ROOT=folder.name)
        self.root.enable(); self.addCleanup(self.root.disable)

    def test_format_is_personal_encrypted_and_not_a_financial_edit(self):
        original = self.add_record(); sheet = self.prepare(); row = sheet.rows.get()
        raw, working = bytes(original.data_ciphertext), bytes(row.data_ciphertext)
        response = self.post("presentation", {"version": 0, "operations": [
            {"kind": "cells", "rows": [str(row.pk)], "columns": ["NOMBRES"],
             "style": {"background": "#FFF2CC", "bold": True, "underline": True}},
        ]}, self.worker)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()["version"], 1)
        own = self.post("query", {"scope": "clean"}, self.worker).json()["presentation"]
        other = self.post("query", {"scope": "clean"}, self.boss).json()["presentation"]
        self.assertTrue(own["cells"][str(row.pk)]["NOMBRES"]["bold"])
        self.assertEqual(other["version"], 0)
        self.assertNotIn(b"#FFF2CC", location(sheet, self.worker).read_bytes())
        sheet.refresh_from_db(); row.refresh_from_db(); original.refresh_from_db()
        self.assertEqual((sheet.revision, row.version), (1, 1))
        self.assertEqual(bytes(original.data_ciphertext), raw)
        self.assertEqual(bytes(row.data_ciphertext), working)
        self.assertFalse(row.edits.exists())

    def test_pagination_and_full_export_keep_stable_row_formats(self):
        for number in range(102):
            self.add_record(f"{number:06}")
        sheet = self.prepare(); last = sheet.rows.order_by("record__source_row").last()
        save_presentation(sheet, self.boss, 0, [
            {"kind": "cells", "rows": [str(last.pk)], "columns": ["NOMBRES"],
             "style": {"background": "#FFF2CC", "color": "#123456", "bold": True, "italic": True, "underline": True, "align": "center", "wrap": True, "size": 16}},
            {"kind": "width", "columns": ["NOMBRES"], "value": 300},
            {"kind": "height", "rows": [str(last.pk)], "value": 64},
        ])
        first = self.post("query", {"scope": "clean", "page": 1}).json()
        second = self.post("query", {"scope": "clean", "page": 2}).json()
        self.assertEqual(first["presentation"]["cells"], {})
        self.assertIn(str(last.pk), second["presentation"]["cells"])
        stream, _ = export_sheet(self.batch, self.boss, "clean", expected_revision=1)
        try:
            book = load_workbook(BytesIO(stream.read()))
            data = book["Principal"]
            self.assertEqual(data.max_row, 103)
            cell = data["C103"]
            self.assertEqual(cell.value, "Persona sintética")
            self.assertEqual(cell.fill.fgColor.rgb[-6:], "FFF2CC")
            self.assertEqual(cell.font.color.rgb[-6:], "123456")
            self.assertTrue(cell.font.bold and cell.font.italic)
            self.assertEqual(cell.font.underline, "single")
            self.assertEqual(cell.font.sz, 12)
            self.assertEqual(cell.alignment.horizontal, "center")
            self.assertTrue(cell.alignment.wrap_text)
            self.assertAlmostEqual(data.column_dimensions["C"].width, 42.14, places=2)
            self.assertEqual(data.row_dimensions[103].height, 48)
            self.assertEqual(data["D103"].value, "000101")
            self.assertEqual(data["I103"].value, 10)
            book.close()
        finally:
            stream.close()

    def test_column_format_is_global_and_problem_export_is_filtered(self):
        self.add_record("GOOD"); self.add_record("BAD", CEDULA_RUC="123")
        sheet = self.prepare()
        save_presentation(sheet, self.boss, 0, [{"kind": "columns", "columns": ["NOMBRES"], "style": {"underline": True}}])
        stream, _ = export_sheet(self.batch, self.boss, "problem", expected_revision=1)
        try:
            book = load_workbook(BytesIO(stream.read()))
            self.assertEqual(book["Problemas"].max_row, 2)
            self.assertEqual(book["Problemas"]["C2"].font.underline, "single")
            book.close()
        finally:
            stream.close()

    def test_stale_save_does_not_overwrite_other_tab(self):
        self.add_record(); sheet = self.prepare()
        save_presentation(sheet, self.boss, 0, [{"kind": "width", "columns": ["NOMBRES"], "value": 400}])
        with self.assertRaises(SheetError) as error:
            save_presentation(sheet, self.boss, 0, [{"kind": "zoom", "value": 120}])
        self.assertEqual(error.exception.code, "presentation_conflict")
        self.assertEqual(read_document(sheet, self.boss)["zoom"], 100)

    def test_invalid_css_foreign_rows_and_massive_operations_rejected_atomically(self):
        self.add_record(); sheet = self.prepare(); row = sheet.rows.get()
        for op in [
            {"kind": "columns", "columns": ["NOMBRES"], "style": {"background": "url(javascript:alert(1))"}},
            {"kind": "columns", "columns": ["NOMBRES"], "style": {"size": 500}},
            {"kind": "cells", "columns": ["NOMBRES"], "rows": ["00000000-0000-0000-0000-000000000000"], "style": {"bold": True}},
            {"kind": "height", "rows": [str(row.pk)], "value": -1},
            {"kind": "width", "columns": ["UNKNOWN"], "value": 200},
        ]:
            self.assertEqual(self.post("presentation", {"version": 0, "operations": [{"kind": "zoom", "value": 120}, op]}).status_code, 400)
        self.assertEqual(read_document(sheet, self.boss)["version"], 0)

    def test_read_only_and_csrf_are_enforced(self):
        self.add_record(); self.prepare()
        viewer = get_user_model().objects.create_user(username="visual-auditor", role="auditor", area="Financiera")
        viewer.user_permissions.add(Permission.objects.get(codename="view_batch"))
        payload = {"version": 0, "operations": [{"kind": "zoom", "value": 110}]}
        self.assertEqual(self.post("presentation", payload, viewer).status_code, 403)
        client = Client(enforce_csrf_checks=True); client.force_login(self.boss)
        self.assertEqual(client.post(f"/api/v1/worksheets/{self.batch.id}/presentation/", data="{}", content_type="application/json").status_code, 403)

    def test_reset_restores_visual_defaults_without_changing_values(self):
        self.add_record(); sheet = self.prepare(); row = sheet.rows.get()
        save_presentation(sheet, self.boss, 0, [
            {"kind": "width", "columns": ["NOMBRES"], "value": 400},
            {"kind": "height", "rows": [str(row.pk)], "value": 80},
            {"kind": "zoom", "value": 125},
        ])
        save_presentation(sheet, self.boss, 1, [{"kind": "reset_sizes"}])
        doc = read_document(sheet, self.boss)
        self.assertEqual(doc["heights"], {})
        self.assertEqual(doc["columns"]["NOMBRES"], {})
        self.assertEqual(doc["zoom"], 100)
        self.assertEqual(values_of(row)["TOTAL_EMISION"], "10.00")
