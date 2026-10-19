"""Regresiones nuevas: columnas de origen, identidades y sólo dos descargas."""
from io import BytesIO
import json
from pathlib import Path
import tempfile
from unittest.mock import patch
from zipfile import ZipFile

from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, override_settings
from openpyxl import Workbook, load_workbook

from apps.batches.crypto import decrypt_json, encrypt_json
from apps.batches.models import Batch, Record
from apps.validation.rules import normalize_header
from apps.validation.services import _record_aad
from apps.validation.tasks import validate_batch
from . import test_worksheet as fixtures
from .columns import CONTRACT_VERSION, columns_for
from .models import ChangeRequest, WorksheetEdit
from .worksheet import FIELDS, decide_change, save_changes, summary, values_of


class ColumnContractTests(TestCase):
    add_record = fixtures.WorksheetTests.add_record
    prepare = fixtures.WorksheetTests.prepare
    post = fixtures.WorksheetTests.post

    def setUp(self):
        self.keys = override_settings(**fixtures.WorksheetTests._overridden_settings)
        self.keys.enable(); self.addCleanup(self.keys.disable)
        fixtures.WorksheetTests.setUp(self)
        folder = tempfile.TemporaryDirectory(prefix="test-column-contract-", dir=Path(__file__).resolve().parent)
        self.addCleanup(folder.cleanup)
        setting = override_settings(ARNOLD_FILE_ROOT=Path(folder.name), ARNOLD_REQUIRE_ANTIVIRUS=False, ARNOLD_LOCAL_QUEUE=False)
        setting.enable(); self.addCleanup(setting.disable)

    @staticmethod
    def upload_file():
        """Una cartera por título. No usa datos ni fixtures de libros reales."""
        book = Workbook(); sheet = book.active; sheet.title = "Cartera"
        labels = ["CIU", "CEDULA_RUC", "NOMBRES", "CLAVE", "CATASTRO", "TITULO", "TITULOS", "TOTAL_EMISION", "TOTAL_INTERES",
                  "TOTAL_COACTIVA", "TOTAL_RECARGO", "MESES_DEUDA", "CORREO", "NOTA MUNICIPAL", "ORDEN"]
        sheet.append(labels)
        for number in range(111):
            key = "REPETIDA" if number in (107, 108) else f"{number:07d}"
            catastro = "" if number == 109 else f"{9000 + number:09d}"
            email = "correo-invalido" if number == 110 else ""
            sheet.append(["000042", "1710034065", "PERSONA SINTETICA", key, catastro, "0007", 7, 10.11, 2.22, 3.33, 4.44, 7, email, "Dato extra sintético", "" if number == 109 else "030"])
        stream = BytesIO(); book.save(stream); book.close()
        return SimpleUploadedFile("cartera-nueva-sintetica.xlsx", stream.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

    def uploaded_pipeline(self):
        self.client.force_login(self.boss)
        uploaded = self.upload_file()
        original_bytes = uploaded.read(); uploaded.seek(0)
        with patch.object(validate_batch, "delay"):
            response = self.client.post("/api/v1/batches/uploads/", {"file": uploaded, "period": "2026-10-01", "schema_name": "carteraMensual",
                       "profile_id": "cartera_mensual_strict", "sheet_name": "Cartera", "header_row": "1"})
        self.assertEqual(response.status_code, 202, response.content)
        result = validate_batch.run(response.json()["run_id"])
        self.assertEqual(result["status"], "succeeded", result)
        self.batch = Batch.objects.get(pk=response.json()["batch_id"])
        from apps.batches.services import batch_file_path
        self.assertEqual(batch_file_path(self.batch).read_bytes(), original_bytes)
        return self.batch.worksheet

    def download(self, scope, revision=1, **extra):
        response = self.post("download", {"scope": scope, "revision": revision, **extra})
        self.assertEqual(response.status_code, 200, getattr(response, "content", b""))
        content = b"".join(response.streaming_content); response.close()
        with ZipFile(BytesIO(content)) as archive:
            self.assertIsNone(archive.testzip())
            self.assertFalse(any("externalLink" in name for name in archive.namelist()))
        book = load_workbook(BytesIO(content)); self.addCleanup(book.close)
        return book, content

    def test_full_upload_validation_partition_query_and_both_complete_downloads(self):
        sheet = self.uploaded_pipeline()
        self.assertEqual(summary(sheet)["clean"], 108)
        self.assertEqual(summary(sheet)["problem"], 3)
        data = self.post("query", {"scope": "clean", "page": 1}).json()
        self.assertEqual(len(data["rows"]), 100)
        self.assertEqual(data["column_contract"], CONTRACT_VERSION)
        all_ids = []
        for scope, count, name in (("clean", 107, "030-SAN CARLOS (LISTA)"), ("problem", 4, "INCONSISTENCIAS")):
            expected = [column["label"] or None for column in self.post("query", {"scope": scope}).json()["columns"]]
            book, _ = self.download(scope, search="0000000", page=1)
            self.assertEqual([cell.value for cell in book[name][1]], expected)
            self.assertEqual(book[name].max_row - 1, count)
            fields = {value: index for index, value in enumerate(expected)}
            rows = list(book[name].iter_rows(min_row=2))
            for row in rows:
                if scope == "problem": self.assertEqual(row[fields["ESTADO_REVISION"]].value, scope)
                self.assertEqual(row[fields["CLAVE"]].data_type, "s")
                if row[fields["CATASTRO"]].value: self.assertEqual(row[fields["CATASTRO"]].data_type, "s")
                self.assertEqual(row[fields["TITULO"]].value, "0007")
                self.assertEqual(row[fields["TOTAL_EMISION"]].data_type, "n")
                self.assertEqual(row[fields["NOTA MUNICIPAL"]].value, "Dato extra sintético")
                self.assertFalse(any(cell.data_type in ("f", "e") for cell in row))
            meta = dict(book["Trazabilidad"].values)
            ids = {key for key in meta if str(key).startswith("Record ")}; all_ids.append(ids)
            self.assertEqual(meta["Filas"], count)
            self.assertEqual(meta["Contrato columnas"], CONTRACT_VERSION)
        self.assertFalse(all_ids[0] & all_ids[1]); self.assertEqual(len(all_ids[0] | all_ids[1]), 111)

    def test_clave_catastro_and_title_counts_never_fill_each_other(self):
        self.add_record("000001", CATASTRO="000099", TITULO="0007")
        self.add_record("000002", CATASTRO="000099", TITULO="0007")
        self.add_record("000003", CATASTRO="", TITULO="0007")
        sheet = self.prepare()
        self.assertEqual((summary(sheet)["clean"], summary(sheet)["problem"]), (3, 0))
        row = sheet.rows.order_by("record__source_row").last()
        self.assertEqual(values_of(row)["CATASTRO"], "")
        self.assertNotEqual(values_of(row)["CLAVE"], values_of(row)["TITULO"])
        self.assertFalse(sheet.rows.filter(duplicate=True).exists())

    def test_independent_protected_edits_approval_history_and_original_preserved(self):
        original = self.add_record("000001", CATASTRO="000099", TITULO="7")
        cipher = bytes(original.data_ciphertext)
        sheet = self.prepare()
        result = save_changes(self.batch, self.worker, [{"id": str(original.pk), "version": 1, "values": {"CATASTRO": "000100"}}], "Expediente sintético")
        self.assertEqual(len(result["pending"]), 1)
        before, _ = self.download("clean")
        values = dict(zip(next(before["030-SAN CARLOS (LISTA)"].values), list(before["030-SAN CARLOS (LISTA)"].values)[1]))
        self.assertEqual((values["CLAVE"], values["CATASTRO"], values["TITULO"]), ("000001", "000099", "7"))
        decide_change(self.batch, self.boss, result["pending"][0], True)
        after, _ = self.download("clean", 2)
        values = dict(zip(next(after["030-SAN CARLOS (LISTA)"].values), list(after["030-SAN CARLOS (LISTA)"].values)[1]))
        self.assertEqual((values["CLAVE"], values["CATASTRO"], values["TITULO"]), ("000001", "000100", "7"))
        original.refresh_from_db(); self.assertEqual(bytes(original.data_ciphertext), cipher)
        row = sheet.rows.get()
        log = decrypt_json(WorksheetEdit.objects.get().detail_ciphertext, associated_data=f"worksheet-edit:{WorksheetEdit.objects.get().pk}")
        self.assertEqual(log["fields"], ["CATASTRO"])
        self.assertEqual(log["old"]["CATASTRO"], "000099")
        self.assertEqual(ChangeRequest.objects.get().status, "approved")

    def test_legacy_lots_are_read_without_mutating_or_reclassifying(self):
        record = self.add_record("000001"); sheet = self.prepare(); row = sheet.rows.get()
        from .worksheet import aad
        old = values_of(row); old.pop("CATASTRO"); old.pop("TITULO")
        row.data_ciphertext = encrypt_json(old, associated_data=aad(row)); row.save(update_fields=("data_ciphertext",))
        ciphertext = bytes(row.data_ciphertext)
        data = self.post("query", {"scope": "clean"}).json()
        self.assertIn("Lote histórico", data["compatibility_notice"])
        self.assertEqual(data["rows"][0]["values"]["CATASTRO"], "0000550044")  # Evidencia RAW, no relleno guardado.
        self.assertEqual(data["rows"][0]["identifier_validation"], "not_executed_legacy")
        self.download("clean")
        row.refresh_from_db(); sheet.refresh_from_db()
        self.assertEqual(bytes(row.data_ciphertext), ciphertext)
        self.assertEqual((row.quality, row.version, sheet.revision), ("clean", 1, 1))

    def test_independent_header_normalization_singular_plural_months(self):
        names = [normalize_header(value, "carteraMensual") for value in ("CLAVE", "CATASTRO", "TITULO", "TITULOS", "TITULOS VENCIDOS", "MESES_DEUDA")]
        self.assertEqual(len(names), len(set(names)))

    def test_third_download_and_readonly_original_edits_are_rejected(self):
        record = self.add_record(); self.prepare()
        response = self.post("save", {"changes": [{"id": str(record.pk), "version": 1,
                             "values": {"source:17:NOTA_MUNICIPAL": "No permitido"}}], "reason": "Prueba sintética"})
        self.assertEqual(response.status_code, 400)
        self.assertEqual(response.json()["error"]["code"], "invalid_fields")
        response = self.post("download", {"scope": "final", "revision": 1})
        self.assertEqual(response.status_code, 400)
        self.assertEqual(response.json()["error"]["code"], "invalid_scope")
        self.client.force_login(self.boss)
        html = self.client.get("/").content.decode()
        self.assertNotIn("downloadFinal", html); self.assertNotIn("downloadKind", html)
        self.assertNotIn("CLAVE / título", html)

    def test_sanitized_current_manifest_is_new_and_has_no_data_rows(self):
        path = Path(__file__).resolve().parents[3] / "docs" / "current_reference_manifest.json"
        manifest = json.loads(path.read_text(encoding="utf-8"))
        self.assertEqual(manifest["sha256"], "c1527fbbca77b4d92f38f420f284de65851b9f3bac68a58d33e9c25109dda29f")
        self.assertEqual(manifest["version"], "current-reference-1")
        for sheet in manifest["sheets"]:
            self.assertNotIn("rows", sheet)
        self.assertTrue(any(s["name"] == "COACTIVA" for s in manifest["sheets"]))
