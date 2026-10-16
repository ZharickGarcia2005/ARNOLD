"""Garantías del flujo: evidencia RAW, clasificación, concurrencia y descarga."""

import base64
from datetime import date, datetime
from io import BytesIO
import json
import uuid

from django.contrib.auth import get_user_model
from django.contrib.auth.models import Permission
from django.test import Client, TestCase, override_settings
from openpyxl import load_workbook

from apps.batches.crypto import encrypt_json
from apps.batches.models import Batch, Record
from apps.validation.contracts import FieldCell, RecordInput
from apps.validation.models import ValidationRun
from apps.validation.services import _record_aad, _record_payload

from .models import ChangeRequest, WorksheetEdit
from .sheet_exports import export_sheet
from .worksheet import FIELDS, SheetError, decide_change, filter_rows, prepare_worksheet, row_payload, save_changes, summary, valid_cedula, values_of


@override_settings(ARNOLD_DATA_ENCRYPTION_KEY=base64.b64encode(b"t" * 32).decode(), ARNOLD_HMAC_KEY="k" * 32)
class WorksheetTests(TestCase):
    def setUp(self):
        users = get_user_model()
        self.boss = users.objects.create_user(username="boss", role="supervisor", area="Financiera", is_superuser=True)
        self.worker = users.objects.create_user(username="worker", area="Financiera")
        self.worker.user_permissions.add(*Permission.objects.filter(codename__in=("view_batch", "change_record", "export_batch")))
        self.batch = Batch.objects.create(display_name="sintetica.xlsx", sha256="c" * 64, period=date(2026, 8, 1),
                                          source_type="cartera_mensual", schema_name="carteraMensual", owner_area="Financiera",
                                          status=Batch.Status.READY, uploaded_by=self.boss)
        self.run = ValidationRun.objects.create(current_batch=self.batch, created_by=self.boss, input_digest="f" * 64,
                                                 engine_version="test", adapter_version="test", status="succeeded")

    def add_record(self, key="000123", **overrides):
        values = dict.fromkeys(FIELDS, "")
        values.update(CIU="000042", CEDULA_RUC="1710034065", NOMBRES="Persona sintética", CLAVE=key,
                      TOTAL_EMISION="10.00", TOTAL_INTERES="2.00", TOTAL_COACTIVA="3.00", TOTAL_RECARGO="4.00", MESES_DEUDA="1")
        values.update(overrides)
        record_id = uuid.uuid4()
        source_row = Record.objects.filter(batch=self.batch).count() + 2
        item = RecordInput(dataset_id=str(self.batch.id), record_id=str(record_id), source_row=source_row,
                           fields={field: FieldCell(value, raw=value) for field, value in values.items()})
        record = Record.objects.create(id=record_id, batch=self.batch, source_row=source_row,
                                      data_ciphertext=encrypt_json(_record_payload(item), associated_data=_record_aad(self.batch.id, record_id, source_row)))
        return record

    def prepare(self):
        self.batch.row_count = self.batch.records.count()
        self.batch.save(update_fields=("row_count",))
        return prepare_worksheet(self.batch)

    def post(self, suffix, data, user=None):
        self.client.force_login(user or self.boss)
        return self.client.post(f"/api/v1/worksheets/{self.batch.id}/{suffix}/", data=json.dumps(data), content_type="application/json")

    def test_partition_preserves_all_rows_and_does_not_deduplicate_by_person(self):
        self.assertTrue(valid_cedula("1710034065"))
        self.add_record("001")
        self.add_record("002")  # misma persona, otro título válido
        self.add_record("DUP")
        self.add_record("dup")
        self.add_record("BAD", CEDULA_RUC="123", NOMBRES="")
        sheet = self.prepare()
        counts = summary(sheet)
        self.assertEqual((counts["total"], counts["clean"], counts["problem"], counts["duplicates"]), (5, 2, 3, 2))
        self.assertEqual(set(filter_rows(sheet, "clean").values_list("record_id", flat=True)).intersection(filter_rows(sheet, "problem").values_list("record_id", flat=True)), set())

    def test_excel_integral_months_and_international_phone_are_valid(self):
        self.add_record(MESES_DEUDA="52.0", TELEF="+593 99-123-4567")
        sheet = self.prepare()
        self.assertEqual(summary(sheet)["clean"], 1)

    def test_edit_is_persistent_and_original_record_remains_unchanged(self):
        original = self.add_record()
        ciphertext = bytes(original.data_ciphertext)
        sheet = self.prepare()
        row = sheet.rows.get()
        result = save_changes(self.batch, self.worker, [{"id": str(row.pk), "version": 1, "values": {"NOMBRES": "Nombre revisado"}}], "Confirmado con expediente")
        row.refresh_from_db(); original.refresh_from_db()
        self.assertEqual(values_of(row)["NOMBRES"], "Nombre revisado")
        self.assertEqual(row.version, 2)
        self.assertEqual(result["summary"]["revision"], 2)
        self.assertEqual(bytes(original.data_ciphertext), ciphertext)
        self.assertNotIn(b"Nombre revisado", bytes(row.data_ciphertext))
        self.assertEqual(WorksheetEdit.objects.get().actor, self.worker)

    def test_invalid_edit_moves_out_of_clean_and_fix_returns_it(self):
        self.add_record(); sheet = self.prepare(); row = sheet.rows.get()
        save_changes(self.batch, self.boss, [{"id": str(row.pk), "version": 1, "values": {"TOTAL_INTERES": "-2"}}], "Revisión")
        row.refresh_from_db(); self.assertEqual(row.quality, "problem")
        save_changes(self.batch, self.boss, [{"id": str(row.pk), "version": 2, "values": {"TOTAL_INTERES": "2.00"}}], "Corregido")
        row.refresh_from_db(); self.assertEqual(row.quality, "clean")

    def test_duplicate_edit_reclassifies_both_sides(self):
        self.add_record("001"); self.add_record("002"); sheet = self.prepare()
        row = sheet.rows.order_by("record__source_row").last()
        save_changes(self.batch, self.boss, [{"id": str(row.pk), "version": 1, "values": {"CLAVE": "001"}}], "Revisión de clave")
        self.assertEqual(sheet.rows.filter(quality="problem").count(), 2)
        save_changes(self.batch, self.boss, [{"id": str(row.pk), "version": 2, "values": {"CLAVE": "002"}}], "Clave confirmada")
        self.assertEqual(sheet.rows.filter(quality="clean").count(), 2)

    def test_stale_multiline_save_is_atomic(self):
        self.add_record("001"); self.add_record("002"); sheet = self.prepare()
        first, second = list(sheet.rows.order_by("record__source_row"))
        with self.assertRaises(SheetError) as raised:
            save_changes(self.batch, self.boss, [
                {"id": str(first.pk), "version": 1, "values": {"NOMBRES": "Cambio parcial"}},
                {"id": str(second.pk), "version": 0, "values": {"NOMBRES": "Obsoleto"}},
            ], "Verificación")
        self.assertEqual(raised.exception.status, 409)
        first.refresh_from_db()
        self.assertEqual(values_of(first)["NOMBRES"], "Persona sintética")
        self.assertFalse(WorksheetEdit.objects.exists())

    def test_protected_change_needs_a_second_user_and_version(self):
        self.add_record(); sheet = self.prepare(); row = sheet.rows.get()
        result = save_changes(self.batch, self.worker, [{"id": str(row.pk), "version": 1, "values": {"TOTAL_EMISION": "20.00"}}], "Respaldo revisado")
        row.refresh_from_db()
        self.assertEqual(values_of(row)["TOTAL_EMISION"], "10.00")
        self.assertEqual(len(result["pending"]), 1)
        repeated = save_changes(self.batch, self.worker, [{"id": str(row.pk), "version": 1, "values": {"TOTAL_EMISION": "20.00"}}], "Reintento de guardado")
        self.assertEqual(repeated["pending"], result["pending"])
        self.assertEqual(ChangeRequest.objects.count(), 1)
        with self.assertRaises(SheetError):
            decide_change(self.batch, self.worker, result["pending"][0], True)
        decide_change(self.batch, self.boss, result["pending"][0], True)
        row.refresh_from_db()
        self.assertEqual(values_of(row)["TOTAL_EMISION"], "20.00")
        self.assertEqual(ChangeRequest.objects.get().decided_by, self.boss)

    def test_api_rejects_cross_area_edits_and_applies_permission_checks(self):
        self.add_record(); sheet = self.prepare(); row = sheet.rows.get()
        other = get_user_model().objects.create_user(username="outside", area="Otra área")
        other.user_permissions.add(*Permission.objects.filter(codename__in=("view_batch", "change_record")))
        response = self.post("save", {"reason": "Prueba", "changes": [{"id": str(row.pk), "version": 1, "values": {"NOMBRES": "No autorizado"}}]}, other)
        self.assertEqual(response.status_code, 404)
        viewer = get_user_model().objects.create_user(username="viewer", role="auditor", area="Financiera")
        viewer.user_permissions.add(Permission.objects.get(codename="view_batch"))
        self.assertEqual(self.post("save", {"reason": "Prueba", "changes": []}, viewer).status_code, 403)
        self.assertEqual(self.post("download", {"scope": "clean", "revision": 1}, viewer).status_code, 403)

    def test_xlsx_exports_all_matching_rows_and_latest_values(self):
        for number in range(102):
            self.add_record(f"{number:06}", FECHA_VENCIMIENTO="2026-08-31" if number == 101 else "")
        sheet = self.prepare()
        row = sheet.rows.order_by("record__source_row").last()
        save_changes(self.batch, self.boss, [{"id": str(row.pk), "version": 1, "values": {"NOMBRES": "Última fila revisada"}}], "Prueba de descarga")
        response = self.post("query", {"scope": "clean", "page": 1})
        self.assertEqual(len(response.json()["rows"]), 100)
        self.assertEqual(response.json()["filtered_total"], 102)
        last_page = self.post("query", {"scope": "clean", "page": 999}).json()
        self.assertEqual(last_page["page"], 2)
        self.assertEqual(len(last_page["rows"]), 2)
        stream, _ = export_sheet(self.batch, self.boss, "clean", expected_revision=2)
        try:
            workbook = load_workbook(BytesIO(stream.read()), data_only=False)
            rows = list(workbook["Principal"].values)
            self.assertEqual(len(rows), 103)
            self.assertEqual(rows[-1][2], "Última fila revisada")
            self.assertEqual(rows[-1][3], "000101")
            self.assertEqual(rows[-1][8], 10)
            self.assertEqual(rows[-1][14], 19)
            self.assertIsInstance(rows[-1][13], datetime)
            self.assertEqual(rows[-1][13].date(), date(2026, 8, 31))
            workbook.close()
        finally:
            stream.close()

    def test_problem_export_contains_only_problems_and_no_formulas(self):
        self.add_record("GOOD"); self.add_record("BAD", NOMBRES="=1+1")
        sheet = self.prepare()
        stream, _ = export_sheet(self.batch, self.boss, "problem", expected_revision=1)
        try:
            workbook = load_workbook(BytesIO(stream.read()), data_only=False)
            self.assertEqual(workbook["Problemas"].max_row, 2)
            self.assertEqual(workbook["Problemas"]["C2"].data_type, "s")
            self.assertTrue(workbook["Problemas"]["C2"].value.startswith("'="))
            self.assertIn("fórmula", workbook["Problemas"]["R2"].value)
            workbook.close()
        finally:
            stream.close()

    def test_stale_export_and_save_return_conflict(self):
        self.add_record(); sheet = self.prepare(); row = sheet.rows.get()
        save_changes(self.batch, self.boss, [{"id": str(row.pk), "version": 1, "values": {"DIRECCION": "Nueva dirección"}}], "Verificado")
        self.assertEqual(self.post("download", {"scope": "clean", "revision": 1}).status_code, 409)
        self.assertEqual(self.post("save", {"changes": [{"id": str(row.pk), "version": 1, "values": {"DIRECCION": "Vieja"}}], "reason": "Prueba"}).status_code, 409)

    def test_search_covers_other_pages_without_exposing_names_in_query(self):
        self.add_record("00001"); self.add_record("00002"); self.prepare()
        response = self.post("query", {"scope": "clean", "search": "00002"}, self.worker)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Cache-Control"], "no-store")
        self.assertEqual(response.json()["filtered_total"], 1)
        self.assertEqual(response.json()["rows"][0]["values"]["CLAVE"], "00002")

    def test_save_endpoint_keeps_csrf_protection(self):
        self.add_record(); self.prepare()
        client = Client(enforce_csrf_checks=True); client.force_login(self.boss)
        response = client.post(f"/api/v1/worksheets/{self.batch.id}/save/", data="{}", content_type="application/json")
        self.assertEqual(response.status_code, 403)

    def test_download_endpoint_returns_complete_http_xlsx_for_each_partition(self):
        for number in range(104):
            self.add_record(f"{number:06}", CEDULA_RUC="123" if number > 100 else "1710034065")
        self.prepare()
        for scope, name, count in (("clean", "Principal", 101), ("problem", "Problemas", 3)):
            response = self.post("download", {"scope": scope, "revision": 1})
            self.assertEqual(response.status_code, 200)
            self.assertEqual(response["Content-Type"], "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
            self.assertIn("attachment;", response["Content-Disposition"])
            self.assertIn(".xlsx", response["Content-Disposition"])
            self.assertEqual(response["X-ARNOLD-Export-Rows"], str(count))
            self.assertEqual(response["X-ARNOLD-Export-Revision"], "1")
            content = b"".join(response.streaming_content)
            response.close()
            self.assertTrue(content.startswith(b"PK\x03\x04"))
            book = load_workbook(BytesIO(content))
            self.assertEqual(book[name].max_row - 1, count)
            self.assertEqual(dict(book["Trazabilidad"].values)["Filas"], count)
            book.close()

    def test_empty_export_is_rejected_instead_of_creating_empty_file(self):
        self.add_record(); self.prepare()
        response = self.post("download", {"scope": "problem", "revision": 1})
        self.assertEqual(response.status_code, 422)
        self.assertEqual(response.json()["error"]["code"], "empty_export")

    def test_original_headers_order_extra_columns_and_selected_duplicates_survive(self):
        self.add_record(); sheet = self.prepare(); row = sheet.rows.get()
        raw = row.record
        from apps.batches.crypto import decrypt_json
        payload = decrypt_json(raw.data_ciphertext, associated_data=_record_aad(self.batch.id, raw.id, raw.source_row))
        payload["fields"].update({
            "NOTA_MUNICIPAL": {"value": "=NO_EJECUTAR()", "raw": "=NO_EJECUTAR()", "kind": "formula"},
            "NOMBRES__1": {"value": "Origen alterno", "raw": "Origen alterno", "kind": "value"},
            "NOMBRES__2": payload["fields"]["NOMBRES"],
        })
        raw.data_ciphertext = encrypt_json(payload, associated_data=_record_aad(self.batch.id, raw.id, raw.source_row)); raw.save()
        headers = ["Nota municipal", "Nombres", "  TOTAL EMISION  ", "Nombres"]
        self.batch.headers_ciphertext = encrypt_json(headers, associated_data=f"batch-headers:{self.batch.id}")
        self.batch.selected_occurrences = {"NOMBRES": 2}; self.batch.save()
        save_changes(self.batch, self.boss, [{"id": str(row.pk), "version": 1, "values": {"NOMBRES": "Nombre aprobado"}}], "Respaldo")
        stream, _ = export_sheet(self.batch, self.boss, "clean", expected_revision=2)
        try:
            book = load_workbook(stream)
            data = book["Principal"]
            self.assertEqual([cell.value for cell in data[1]][:4], headers)
            self.assertEqual([cell.value for cell in data[2]][:4], ["'=NO_EJECUTAR()", "Origen alterno", 10, "Nombre aprobado"])
            self.assertEqual(data["C2"].number_format, "#,##0.00")
            self.assertEqual(data.auto_filter.ref, f"A1:{data.cell(1, data.max_column).column_letter}2")
            self.assertEqual(data["A2"].data_type, "s")
            book.close()
        finally:
            stream.close()

    def test_reason_cannot_be_bypassed_and_pending_money_is_not_exported(self):
        self.add_record(); sheet = self.prepare(); row = sheet.rows.get()
        changes = [{"id": str(row.pk), "version": 1, "values": {"TOTAL_EMISION": "20.00"}}]
        response = self.post("save", {"changes": changes, "reason": ""}, self.worker)
        self.assertEqual(response.status_code, 400)
        self.assertFalse(WorksheetEdit.objects.exists()); self.assertFalse(ChangeRequest.objects.exists())
        result = self.post("save", {"changes": changes, "reason": "Documentado"}, self.worker).json()
        self.assertEqual(len(result["pending"]), 1)
        response = self.post("download", {"scope": "clean", "revision": result["summary"]["revision"]}, self.worker)
        book = load_workbook(BytesIO(b"".join(response.streaming_content)))
        response.close(); self.assertEqual(book["Principal"]["I2"].value, 10); book.close()
