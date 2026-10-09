"""Build privacy-safe source profiles and cross-file validation evidence for ARNOLD.

Only aggregate counts, schemas, and rule evidence are printed. Row-level values and
identifiers remain in memory and are never written to the project.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import unicodedata
from collections import Counter, defaultdict
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


def norm(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    text = unicodedata.normalize("NFKD", str(value))
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    return re.sub(r"\s+", " ", text).strip().upper()


def unique_headers(values: list[Any]) -> tuple[list[str], list[str]]:
    seen: Counter[str] = Counter()
    result: list[str] = []
    duplicates: list[str] = []
    for idx, value in enumerate(values, start=1):
        base = norm(value) or f"COL_{idx}"
        seen[base] += 1
        if seen[base] > 1:
            duplicates.append(base)
            result.append(f"{base}__{seen[base]}")
        else:
            result.append(base)
    return result, sorted(set(duplicates))


def row_hash(values: list[str]) -> str:
    return hashlib.sha256("\x1f".join(values).encode("utf-8")).hexdigest()


def load_tables(path: Path, sheet_names: list[str]) -> dict[str, dict[str, Any]]:
    workbook = load_workbook(path, read_only=True, data_only=True, keep_links=False)
    tables: dict[str, dict[str, Any]] = {}
    for sheet_name in sheet_names:
        ws = workbook[sheet_name]
        iterator = ws.iter_rows(values_only=True)
        try:
            first = list(next(iterator))
        except StopIteration:
            tables[sheet_name] = {"headers": [], "duplicates": [], "rows": [], "error_cells": {}}
            continue
        headers, duplicate_headers = unique_headers(first)
        rows: list[dict[str, str]] = []
        errors: Counter[str] = Counter()
        exact_hashes: Counter[str] = Counter()
        for raw in iterator:
            values = [norm(value) for value in raw]
            if not any(values):
                continue
            if len(values) < len(headers):
                values.extend([""] * (len(headers) - len(values)))
            elif len(values) > len(headers):
                values = values[: len(headers)]
            for value in values:
                if value.startswith("#"):
                    errors[value] += 1
            exact_hashes[row_hash(values)] += 1
            rows.append(dict(zip(headers, values)))
        tables[sheet_name] = {
            "headers": headers,
            "duplicate_headers": duplicate_headers,
            "rows": rows,
            "error_cells": dict(errors),
            "exact_duplicate_extra_rows": sum(count - 1 for count in exact_hashes.values() if count > 1),
        }
    workbook.close()
    return tables


def make_key(row: dict[str, str], fields: tuple[str, ...]) -> str:
    values = [row.get(field, "") for field in fields]
    return "\x1e".join(values) if all(values) else ""


def key_profile(rows: list[dict[str, str]], fields: tuple[str, ...], identity_field: str) -> dict[str, int]:
    grouped: defaultdict[str, list[dict[str, str]]] = defaultdict(list)
    for row in rows:
        key = make_key(row, fields)
        if key:
            grouped[key].append(row)
    duplicates = {key: values for key, values in grouped.items() if len(values) > 1}
    return {
        "nonblank_rows": sum(len(values) for values in grouped.values()),
        "distinct": len(grouped),
        "duplicate_groups": len(duplicates),
        "duplicate_extra_rows": sum(len(values) - 1 for values in duplicates.values()),
        "groups_with_multiple_identities": sum(
            len({row.get(identity_field, "") for row in values if row.get(identity_field)}) > 1
            for values in duplicates.values()
        ),
    }


def numeric_profile(rows: list[dict[str, str]], field: str) -> dict[str, int]:
    counts: Counter[str] = Counter()
    for row in rows:
        value = row.get(field, "")
        if not value:
            counts["blank"] += 1
            continue
        candidate = value.replace("$", "").replace(" ", "")
        if "," in candidate and "." in candidate:
            candidate = candidate.replace(",", "")
        elif "," in candidate:
            candidate = candidate.replace(",", ".")
        try:
            number = Decimal(candidate)
        except (InvalidOperation, ValueError):
            counts["nonnumeric"] += 1
            continue
        if number < 0:
            counts["negative"] += 1
        elif number == 0:
            counts["zero"] += 1
        else:
            counts["positive"] += 1
    return dict(counts)


def table_profile(table: dict[str, Any], key_candidates: list[tuple[str, ...]], identity_field: str, amount_fields: list[str]) -> dict[str, Any]:
    rows = table["rows"]
    return {
        "row_count": len(rows),
        "column_count": len(table["headers"]),
        "headers": table["headers"],
        "duplicate_headers": table["duplicate_headers"],
        "exact_duplicate_extra_rows": table["exact_duplicate_extra_rows"],
        "error_cells": table["error_cells"],
        "keys": {
            "+".join(fields): key_profile(rows, fields, identity_field)
            for fields in key_candidates
            if all(field in table["headers"] for field in fields)
        },
        "amount_quality": {
            field: numeric_profile(rows, field)
            for field in amount_fields
            if field in table["headers"]
        },
        "missing_required": {
            field: sum(not row.get(field, "") for row in rows)
            for field in (identity_field, "CIU", "CLAVE", "EXP", "ID PLAN")
            if field in table["headers"]
        },
    }


def compare_tables(
    first: dict[str, Any],
    second: dict[str, Any],
    key_fields: tuple[str, ...],
    compare_fields: list[tuple[str, str]],
) -> dict[str, Any]:
    first_map: defaultdict[str, list[dict[str, str]]] = defaultdict(list)
    second_map: defaultdict[str, list[dict[str, str]]] = defaultdict(list)
    for row in first["rows"]:
        key = make_key(row, key_fields)
        if key:
            first_map[key].append(row)
    for row in second["rows"]:
        key = make_key(row, key_fields)
        if key:
            second_map[key].append(row)
    common = set(first_map) & set(second_map)
    changes: dict[str, int] = {}
    for first_field, second_field in compare_fields:
        changed = 0
        for key in common:
            first_values = {row.get(first_field, "") for row in first_map[key]}
            second_values = {row.get(second_field, "") for row in second_map[key]}
            if first_values != second_values:
                changed += 1
        changes[f"{first_field}->{second_field}"] = changed
    return {
        "key": "+".join(key_fields),
        "first_distinct_keys": len(first_map),
        "second_distinct_keys": len(second_map),
        "common_keys": len(common),
        "only_first": len(set(first_map) - set(second_map)),
        "only_second": len(set(second_map) - set(first_map)),
        "changed_common_key_groups": changes,
    }


def exact_overlap(first: dict[str, Any], second: dict[str, Any], fields: list[str]) -> dict[str, int]:
    first_hashes: Counter[str] = Counter()
    second_hashes: Counter[str] = Counter()
    for row in first["rows"]:
        first_hashes[row_hash([row.get(field, "") for field in fields])] += 1
    for row in second["rows"]:
        second_hashes[row_hash([row.get(field, "") for field in fields])] += 1
    overlap = sum(min(count, second_hashes.get(key, 0)) for key, count in first_hashes.items())
    return {
        "first_rows": sum(first_hashes.values()),
        "second_rows": sum(second_hashes.values()),
        "exact_matching_rows": overlap,
        "only_first_rows": sum(first_hashes.values()) - overlap,
        "only_second_rows": sum(second_hashes.values()) - overlap,
    }


def classify_novelty(value: str) -> str:
    if not value:
        return "blank"
    rules = [
        ("identity", ("CEDULA", "RUC", "IDENTIFIC")),
        ("duplicate", ("DUPLIC", "REPET")),
        ("payment", ("PAG", "CANCEL", "ABON", "SALD")),
        ("agreement", ("CONVEN",)),
        ("name", ("NOMBRE",)),
        ("address_or_location", ("DIRECCION", "PARROQUIA", "SECTOR", "PREDIO")),
        ("amount", ("MONTO", "VALOR", "INTERES", "EMISION", "RECARGO")),
        ("missing_or_not_found", ("NO EXIST", "NO CONSTA", "NO ENCONTR", "SIN DAT")),
    ]
    for label, words in rules:
        if any(word in value for word in words):
            return label
    return "other_nonblank"


def order_union(order_tables: dict[str, dict[str, Any]], names: list[str]) -> dict[str, Any]:
    rows: list[dict[str, str]] = []
    for name in names:
        for row in order_tables[name]["rows"]:
            mapped = dict(row)
            mapped["CEDULA_RUC"] = row.get("CEDULA", "")
            mapped["PREDIO_MUN"] = row.get("PREDIO MUNICIPAL", "")
            mapped["VALOR_TOTAL"] = row.get("TOTAL", "")
            mapped["COACTIVA"] = row.get("RECARGO", "")
            mapped["SOURCE_SHEET"] = name
            rows.append(mapped)
    headers = sorted({key for row in rows for key in row})
    return {"headers": headers, "duplicate_headers": [], "rows": rows, "error_cells": {}, "exact_duplicate_extra_rows": 0}


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--july", type=Path, required=True)
    parser.add_argument("--august", type=Path, required=True)
    parser.add_argument("--legacy", type=Path, required=True)
    parser.add_argument("--orders", type=Path, required=True)
    args = parser.parse_args()

    july = load_tables(args.july, ["Sheet 1", "UNIF-SHEET 1", "SHEET 1-SIN COACTIVA", "Hoja 2"])
    august = load_tables(args.august, ["Sheet 1", "CAT ACTUALIZADO", "SSheet 1", "SHEET-UNIFICADO", "Copia de SHEET-UNIFICADO"])
    legacy = load_tables(args.legacy, ["Sheet 1"])
    order_names = ["030-SAN CARLOS (LISTA)", "031-(24 DE MAYO) LISTA", "032-SAN CRISTOBAL(LISTA)", "033-GUAYACAN", "034-NICOLAS ALTA"]
    orders = load_tables(args.orders, order_names + ["INCONSISTENCIAS"])
    combined_orders = order_union(orders, order_names)

    raw_keys = [("CIU",), ("CLAVE",), ("CEDULA_RUC",), ("CIU", "CLAVE"), ("CEDULA_RUC", "CIU", "CLAVE")]
    unified_keys = [("EXP",), ("TP",), ("CIU",), ("CLAVE",), ("ID PLAN",), ("CEDULA_RUC", "CIU"), ("CEDULA_RUC", "CLAVE"), ("CEDULA_RUC", "CIU", "CLAVE")]
    raw_amounts = ["TOTAL_EMISION", "TOTAL_INTERES", "TOTAL_COACTIVA", "TOTAL_RECARGO", "MESES_DEUDA"]
    unified_amounts = ["EMISION", "INTERES", "SUBTOTAL", "COACTIVA", "TOTAL_RECARGO", "VALOR_TOTAL", "TITULOS"]

    profiles = {
        "july_raw": table_profile(july["Sheet 1"], raw_keys, "CEDULA_RUC", raw_amounts),
        "july_unified": table_profile(july["UNIF-SHEET 1"], unified_keys, "CEDULA_RUC", unified_amounts),
        "july_without_coactiva": table_profile(july["SHEET 1-SIN COACTIVA"], unified_keys, "CEDULA_RUC", unified_amounts),
        "august_raw": table_profile(august["Sheet 1"], raw_keys, "CEDULA_RUC", raw_amounts),
        "august_cleaned_raw": table_profile(august["SSheet 1"], raw_keys, "CEDULA_RUC", raw_amounts + ["T_EMISION", "T_INTERES", "T_COACTIVA", "T_RECARGO", "MES_DEUDA"]),
        "august_unified": table_profile(august["SHEET-UNIFICADO"], unified_keys, "CEDULA_RUC", unified_amounts),
        "august_catalog": table_profile(august["CAT ACTUALIZADO"], [("RUC_CEDULA",), ("CLAVE_NUEV",), ("RUC_CEDULA", "CLAVE_NUEV")], "RUC_CEDULA", []),
        "legacy_july_coactiva": table_profile(legacy["Sheet 1"], raw_keys, "CEDULA_RUC", raw_amounts),
        "preliminary_orders_030_034": table_profile(combined_orders, unified_keys, "CEDULA_RUC", unified_amounts),
        "order_inconsistencies": table_profile(orders["INCONSISTENCIAS"], [("EXP",), ("CEDULA", "CIU"), ("CEDULA", "CLAVE")], "CEDULA", ["EMISION", "INTERES", "RECARGO", "TOTAL"]),
    }

    comparisons = {
        "legacy_vs_august_raw": exact_overlap(legacy["Sheet 1"], august["Sheet 1"], legacy["Sheet 1"]["headers"]),
        "august_unified_copy": exact_overlap(august["SHEET-UNIFICADO"], august["Copia de SHEET-UNIFICADO"], august["SHEET-UNIFICADO"]["headers"]),
        "july_unified_vs_without_coactiva": compare_tables(
            july["UNIF-SHEET 1"], july["SHEET 1-SIN COACTIVA"], ("CEDULA_RUC", "CIU", "CLAVE"),
            [(field, field) for field in ("EXP", "EMISION", "INTERES", "COACTIVA", "VALOR_TOTAL")],
        ),
        "july_vs_august_unified": compare_tables(
            july["UNIF-SHEET 1"], august["SHEET-UNIFICADO"], ("CEDULA_RUC", "CIU", "CLAVE"),
            [(field, field) for field in ("EXP", "TP", "EMISION", "INTERES", "COACTIVA", "TOTAL_RECARGO", "VALOR_TOTAL", "TITULOS")],
        ),
        "august_to_preliminary_orders": compare_tables(
            august["SHEET-UNIFICADO"], combined_orders, ("CEDULA_RUC", "CIU", "CLAVE"),
            [("EXP", "EXP"), ("EMISION", "EMISION"), ("INTERES", "INTERES"), ("COACTIVA", "COACTIVA"), ("VALOR_TOTAL", "VALOR_TOTAL")],
        ),
    }

    novelty_counts = Counter(classify_novelty(row.get("NOVEDAD", "")) for row in orders["INCONSISTENCIAS"]["rows"])
    order_rows = {name: len(orders[name]["rows"]) for name in order_names}
    output = {
        "sources": {
            "july": args.july.name,
            "august": args.august.name,
            "legacy": "CARTERA_VENCIDA_julio_coactiva.xls",
            "orders": args.orders.name,
        },
        "profiles": profiles,
        "comparisons": comparisons,
        "preliminary_order_rows": order_rows,
        "inconsistency_novelty_categories": dict(novelty_counts),
    }
    print(json.dumps(output, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
