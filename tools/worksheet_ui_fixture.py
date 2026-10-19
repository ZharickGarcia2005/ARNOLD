"""Aplicación REAL Django con base aislada y un lote exclusivamente sintético.

Sin --verify-only: http://127.0.0.1:8011/ (boss / Arnold-Visual-2026!).
Con --verify-only: genera las dos descargas usando los endpoints reales.
No lee, migra, reemplaza ni comparte la base operativa del servidor 8000.
"""
from datetime import date
from decimal import Decimal
from io import BytesIO
import json
import os
from pathlib import Path
import sys
import tempfile

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "backend"))
os.environ.setdefault("DJANGO_SETTINGS_MODULE", "core.settings.development")
import django
django.setup()
from django.conf import settings
from django.core.management import call_command
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import Client
from django.test.runner import DiscoverRunner
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from apps.workflow import test_column_contract as fixtures
from apps.workflow.columns import SHEETS, CONTRACT_VERSION
from apps.workflow.order_projection import presentation_counts

class VisualCase(fixtures.ColumnContractTests):
    @staticmethod
    def upload_file():
        upload = fixtures.ColumnContractTests.upload_file()
        book = load_workbook(BytesIO(upload.read()))
        sheet = book.active
        extra = ["EXP", "PREDIO MUNICIPAL", "PARROQUIA", "PARROQ", "ZONA", "SECTOR", "MZ", "PREDIO", "ID PLAN", "DIGITADOR", "FECHA", "HORA", "NOVEDAD", "UNIDAD", "FE", "REGISTRADO", "TOTAL", "TOTAL EN LETRA"]
        first = sheet.max_column + 1
        for c, label in enumerate(extra, first): sheet.cell(1, c, label)
        order_column = next(c.column for c in sheet[1] if c.value == "ORDEN")
        codes = ("036", "037", "038", "039", "040", "031", "035")
        for number in range(111):
            row = number + 2
            if 100 <= number <= 106: sheet.cell(row, order_column, codes[number - 100])
            values = [f"EXP-{number:04}", f"PREDIO-{number:05}", "PARROQUIA SINTETICA", "01", "02", "03", "04", "05", "00001", "OPERADOR DE PRUEBA", date(2026, 10, 1), "09:30:00", "Dato ficticio", "UNIDAD FICTICIA", "FE SINTETICO", "SI", 20.10, "TOTAL DE ORIGEN FICTICIO"]
            for c, value in enumerate(values, first): sheet.cell(row, c, value)
        output = BytesIO(); book.save(output); book.close()
        return SimpleUploadedFile("PRUEBA-SINTETICA-ORDENES.xlsx", output.getvalue(), content_type=upload.content_type)

def verify(case):
    sheet = case.batch.worksheet
    row = sheet.rows.order_by("record__source_row").first()
    original = bytes(row.record.data_ciphertext)
    saved = case.post("save", {"reason":"Nombre confirmado en expediente ficticio", "changes":[{"id":str(row.pk), "version":1, "values":{"NOMBRES":"PERSONA SINTETICA REVISADA"}}]}, case.worker)
    assert saved.status_code == 200
    pending = case.post("save", {"reason":"Clave contrastada con respaldo ficticio", "changes":[{"id":str(row.pk), "version":2, "values":{"CLAVE":"000099999"}}]}, case.worker)
    assert pending.status_code == 200
    approved = case.post(f"requests/{pending.json()['pending'][0]}/decide", {"approve":True}, case.boss)
    assert approved.status_code == 200
    row.record.refresh_from_db(); assert bytes(row.record.data_ciphertext) == original
    output = ROOT / "outputs" / "order-presentation"
    output.mkdir(parents=True, exist_ok=True)
    report, ids = {}, []
    for scope, count, filename in (("clean",107,"principal-por-orden-sintetico.xlsx"),("problem",4,"inconsistencias-sinteticas.xlsx")):
        book, content = case.download(scope, 3, search="0000000", page=1, order="030")
        names = [definition["name"] for code, definition in SHEETS.items() if (code == "issues") == (scope == "problem")]
        assert book.sheetnames == names + ["Trazabilidad"]
        written = 0; totals = dict.fromkeys(("EMISION","INTERES","TOTAL_COACTIVA","RECARGO"), Decimal(0))
        for name in names:
            code = next(code for code in SHEETS if SHEETS[code]["name"] == name)
            data = case.post("query", {"order":code}).json()
            ws = book[name]; labels = [c.value or "" for c in ws[1]]
            assert labels == [c["label"] for c in data["columns"]]
            for c, column in enumerate(data["columns"], 1):
                assert ws.column_dimensions[get_column_letter(c)].hidden == column["hidden"]
            fields = {label:i for i,label in enumerate(labels)}
            written += ws.max_row - 1
            for cells in ws.iter_rows(min_row=2):
                assert isinstance(cells[fields["CLAVE"]].value, str)
                assert not any(c.data_type in ("f","e") for c in cells)
                for field in totals: totals[field] += Decimal(str(cells[fields[field]].value or 0))
            assert ws.freeze_panes == "A2"
        assert written == count
        assert sum(totals.values()) == Decimal("20.10") * count
        meta = dict(book["Trazabilidad"].values); assert meta["Filas"] == count
        ids.append({key for key in meta if str(key).startswith("Record ")})
        (output / filename).write_bytes(content)
        report[scope] = {"rows":written, "sheets":names, "totals":{k:str(v) for k,v in totals.items()}, "contract":CONTRACT_VERSION}
    assert not ids[0] & ids[1] and len(ids[0] | ids[1]) == 111
    result = {"synthetic_only":True, "disjoint":True, "original_unchanged":True, "revision":3, "downloads":report}
    (output / "verification.json").write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(result, ensure_ascii=False, indent=2), flush=True)

def main():
    if settings.DATABASES["default"]["ENGINE"] != "django.db.backends.sqlite3": raise RuntimeError("Sólo SQLite de pruebas.")
    root = ROOT / "outputs"; root.mkdir(exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="order-ui-test-", dir=root) as directory:
        settings.DATABASES["default"].setdefault("TEST", {})["NAME"] = str(Path(directory) / "test.sqlite3")
        runner = DiscoverRunner(verbosity=0, interactive=False)
        runner.setup_test_environment(); old_config = runner.setup_databases()
        # Los puertos no aíslan cookies: no compartir las de la aplicación 8000.
        settings.SESSION_COOKIE_NAME = "arnold_order_ui_session"
        settings.CSRF_COOKIE_NAME = "arnold_order_ui_csrf"
        case = VisualCase(methodName="test_full_upload_validation_partition_query_and_both_complete_downloads")
        case.client = Client()
        try:
            case.setUp(); case.uploaded_pipeline()
            if "--verify-only" in sys.argv: verify(case); return
            case.boss.set_password("Arnold-Visual-2026!"); case.boss.save(update_fields=("password",))
            settings.DEBUG = True
            print("Servidor Django aislado: http://127.0.0.1:8011/ · sólo datos ficticios", flush=True)
            call_command("runserver", "127.0.0.1:8011", use_reloader=False)
        finally:
            case.doCleanups(); runner.teardown_databases(old_config); runner.teardown_test_environment()

if __name__ == "__main__": main()
