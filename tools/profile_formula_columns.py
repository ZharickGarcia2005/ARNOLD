"""Count formulas and cached errors by column without exposing formula text or values."""

from __future__ import annotations

import argparse
import json
import re
import unicodedata
from collections import Counter
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


def norm(value: Any) -> str:
    if value is None:
        return ""
    text = unicodedata.normalize("NFKD", str(value))
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    return re.sub(r"\s+", " ", text).strip().upper()


def unique_headers(values: list[Any]) -> list[str]:
    seen: Counter[str] = Counter()
    headers: list[str] = []
    for idx, value in enumerate(values, start=1):
        base = norm(value) or f"COL_{idx}"
        seen[base] += 1
        headers.append(base if seen[base] == 1 else f"{base}__{seen[base]}")
    return headers


def profile(path: Path, sheet_names: list[str]) -> dict[str, Any]:
    formula_book = load_workbook(path, read_only=True, data_only=False, keep_links=False)
    value_book = load_workbook(path, read_only=True, data_only=True, keep_links=False)
    result: dict[str, Any] = {"file": path.name, "sheets": {}}
    for sheet_name in sheet_names:
        fws = formula_book[sheet_name]
        vws = value_book[sheet_name]
        fiter = fws.iter_rows(values_only=True)
        viter = vws.iter_rows(values_only=True)
        try:
            headers = unique_headers(list(next(fiter)))
            next(viter)
        except StopIteration:
            result["sheets"][sheet_name] = {}
            continue
        stats = {header: Counter() for header in headers}
        for frow, vrow in zip(fiter, viter):
            for idx, header in enumerate(headers):
                formula_value = frow[idx] if idx < len(frow) else None
                cached_value = vrow[idx] if idx < len(vrow) else None
                if isinstance(formula_value, str) and formula_value.startswith("="):
                    stats[header]["formula"] += 1
                    if cached_value is None or cached_value == "":
                        stats[header]["formula_cached_blank"] += 1
                    elif isinstance(cached_value, str) and cached_value.startswith("#"):
                        stats[header]["formula_cached_error"] += 1
                    else:
                        stats[header]["formula_cached_value"] += 1
                elif formula_value is None or formula_value == "":
                    stats[header]["blank_literal"] += 1
                else:
                    stats[header]["literal"] += 1
        result["sheets"][sheet_name] = {
            header: dict(counts)
            for header, counts in stats.items()
            if counts.get("formula") or counts.get("formula_cached_error")
        }
    formula_book.close()
    value_book.close()
    return result


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("path", type=Path)
    parser.add_argument("sheets", nargs="+")
    args = parser.parse_args()
    print(json.dumps(profile(args.path, args.sheets), ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
