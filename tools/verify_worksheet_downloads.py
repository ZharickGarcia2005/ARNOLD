"""Verifica las dos descargas por orden sin escribir ni modificar los XLSX."""
import argparse
from decimal import Decimal
import json
from pathlib import Path
from zipfile import ZipFile
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
MANIFEST = json.loads((ROOT / "backend/apps/workflow/order_manifest.json").read_text(encoding="utf-8"))

def inspect_download(path, scope, expected, *, revision, expected_headers=None):
    with ZipFile(path) as archive:
        assert archive.testzip() is None, "ZIP dañado"
        assert not any("externalLink" in name for name in archive.namelist())
    templates = [s for s in MANIFEST["sheets"] if (s["id"] == "issues") == (scope == "problem")]
    book = load_workbook(path, data_only=False)
    try:
        assert book.sheetnames[:len(templates)] == [s["name"] for s in templates]
        assert all(name.startswith("Trazabilidad") for name in book.sheetnames[len(templates):])
        count = 0; totals = dict.fromkeys(("EMISION", "INTERES", "RECARGO", "TOTAL_COACTIVA"), Decimal(0))
        for template in templates:
            sheet = book[template["name"]]
            headers = [cell.value or "" for cell in sheet[1]]
            assert headers[:len(template["columns"])] == [col["label"] for col in template["columns"]]
            for index, col in enumerate(template["columns"], 1):
                assert sheet.column_dimensions[get_column_letter(index)].hidden == col["hidden"]
            fields = {value: i for i, value in enumerate(headers)}
            for row in sheet.iter_rows(min_row=2):
                count += 1
                for field in ("CLAVE", "CIU", "CEDULA"):
                    cell = row[fields[field]]
                    assert cell.value is None or isinstance(cell.value, str)
                    assert cell.number_format == "@"
                assert not any(cell.data_type in ("f", "e") for cell in row)
                for field in totals:
                    if field in fields and row[fields[field]].data_type == "n" and row[fields[field]].value is not None:
                        totals[field] += Decimal(str(row[fields[field]].value))
            assert sheet.freeze_panes == "A2"
        meta = {}
        for name in book.sheetnames[len(templates):]: meta.update(dict(book[name].values))
        assert count == expected == meta["Filas"]
        assert meta["Conjunto"] == scope and meta["Revisión"] == revision
        ids = {key for key in meta if str(key).startswith("Record ")}
        assert len(ids) == count
        return {"file":str(Path(path).resolve()), "rows":count, "sheets":[s["name"] for s in templates],
                "financial_totals":{k:str(v) for k,v in totals.items()}, "revision":revision, "xlsx_crc":"ok"}, ids
    finally: book.close()

def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("clean"); parser.add_argument("problems")
    parser.add_argument("--clean-count", type=int, required=True)
    parser.add_argument("--problem-count", type=int, required=True)
    parser.add_argument("--revision", type=int, required=True)
    args = parser.parse_args()
    clean, first = inspect_download(args.clean, "clean", args.clean_count, revision=args.revision)
    problems, second = inspect_download(args.problems, "problem", args.problem_count, revision=args.revision)
    assert not first & second
    print(json.dumps({"clean":clean,"problems":problems,"disjoint":True,"total":len(first | second)}, ensure_ascii=False, indent=2))

if __name__ == "__main__": main()
